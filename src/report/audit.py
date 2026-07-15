"""On-demand audit workbook — `poc3 report audit` (issue #186 Option D).

The Audit Trail is QA/methodology-debug material (its sheet-guide
audience was always "Engineers"), not an analyst deliverable. Option D
removes it from the analyst workbooks; this module keeps the full
surface ONE command away:

    poc3 report audit --state TX [--lens spine]

writes ``data/out/{state}_audit{_spine}.xlsx`` — Readme + the full
Audit Trail (EVERY row, documented or not: the audit surface must never
shrink) + Legend. Columns stay `workbook_spec.audit_trail_columns(lens)`
(the composed Details-prefix + fact/span/dim/NACHOS/review blocks), and
``Row #`` stays the shared canonical-sort address, so an audit row still
cross-references the analyst workbook's Details row by number.
"""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Literal, cast

from openpyxl import Workbook

from src.report.analyst import (
    _element_row_contexts,
    _load_inputs,
    _load_scores_sidecar,
    _row_number_index,
    _write_legend_sheet,
    _write_readme_sheet,
)
from src.report.workbook_render import _finalize_workbook, render_table_sheet
from src.report.workbook_spec import AUDIT_TRAIL_SHEET
from src.utils.paths import out_dir

logger = logging.getLogger(__name__)

# Module constant baked at import (test monkeypatch convention). Loads
# go through `analyst._load_inputs` / `_load_scores_sidecar`, which
# resolve against analyst's own `_OUT_DIR`/`_SPINE_DIR` constants — the
# same patch points every analyst test already uses.
_OUT_DIR = out_dir()


def run(
    *,
    state: str,
    lens: str = "source",
    out: Path | None = None,
) -> Path:
    """Write ``{state}_audit{_spine}.xlsx``. Returns the output path.

    Renders with or without a scores sidecar — an unscored audit (facts
    absent, dims blank) is still an honest audit surface.
    """
    out_base = out or _OUT_DIR
    out_base.mkdir(parents=True, exist_ok=True)
    st = state.upper()
    suffix = "_spine" if lens == "spine" else ""

    si = _load_inputs(st, lens=lens)
    scores = _load_scores_sidecar(st, lens, out_base)

    # Recommendations join — the Audit Trail's `AI: Recommendations`
    # column had silently regressed to always-blank when the sheet moved
    # off the analyst workbook (Seq 4 PR C built contexts without a recs
    # lookup; issue #213 item 1). Same optional posture as analyst.run:
    # an absent recs sidecar renders the column empty, honestly.
    recs_lookup: dict[str, list[dict]] | None = None
    rec_path = out_base / f"{st.lower()}_recommendations_{lens}.json"
    if rec_path.exists():
        from src.report.analyst import _recs_lookup
        from src.report.recommendations import sheet_rows_for_workbook

        try:
            recs_lookup = _recs_lookup(
                {st: sheet_rows_for_workbook(
                    st,
                    cast("Literal['source', 'spine']", lens),
                    base=out_base,
                )}
            )
        except FileNotFoundError:
            recs_lookup = None

    wb = Workbook()
    readme_ws = wb.active
    readme_ws.title = "Readme"
    row_numbers = _row_number_index([si])
    contexts = _element_row_contexts(
        [si],
        {st: scores} if scores else None,
        row_numbers,
        recs_lookup=recs_lookup,
    )
    render_table_sheet(
        wb.create_sheet("Audit Trail"), AUDIT_TRAIL_SHEET, contexts, lens
    )
    _write_legend_sheet(wb.create_sheet("Legend"), lens)
    _write_readme_sheet(readme_ws, wb)
    _finalize_workbook(wb, lens)

    path = out_base / f"{st.lower()}_audit{suffix}.xlsx"
    wb.save(path)
    logger.info(
        "audit[%s]: wrote %s (%d rows%s)",
        lens, path, len(contexts),
        "" if scores else "; no scores sidecar — fact columns blank",
    )
    return path

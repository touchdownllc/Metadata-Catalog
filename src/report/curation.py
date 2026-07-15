"""Analyst-edit round-trip — the curation sidecar (issue #186 Option C).

``poc3 review ingest <workbook.xlsx>`` reads the analyst-input band from
a generated analyst workbook back into a per-state curation sidecar at
``data/curation/{state}.json``. Every ``report analyst`` regeneration
re-applies the stored values into the band, so regenerating a workbook
stops destroying analyst work.

The curation sidecar is the ONLY non-regenerable artifact in the repo —
it is human input — so unlike everything under ``data/out/`` it is
COMMITTED to git (``data/curation/`` is not gitignored, deliberately).

Design contract (``next-phase/scoring-pipeline-design.md`` §8.4):
overrides are curation rows, never edits to engine output. Re-scores and
regenerations never touch them; the override displays alongside the
engine score; disagreement between the two is itself a review signal
(see ``override_disagreements`` — surfaced on the Review Queue sheet).

Merge semantics: per-column newest-wins with bounded history. A blank
workbook cell is "no input", NEVER a delete — analysts filter rows, and
a filtered-out row must not lose its stored comments on re-ingest.
There is no delete mechanism in v1; the ingest report says so.

Keying: ``{STATE}|{entity}|{element_name}`` — exactly the score-sidecar
``analyst._record_key`` shape, lens-agnostic (workflow judgments are
about the element, not a lens). Each captured value records the lens of
the workbook it came from as provenance, so lens-sensitive logic
(override disagreement) can gate on it.
"""

from __future__ import annotations

import json
import logging
import re
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from src.report.workbook_spec import (
    COMMITMENT_TRACKER_COLUMNS,
    CellValue,
    DETAILS_COLUMNS,
)
from src.utils.paths import out_dir, project_root

logger = logging.getLogger(__name__)

# v2 (issue #248 Part B): entries may carry an optional `adjudication`
# block (team-consensus score, adjusted axis only) alongside `values` /
# `history`. Additive — v1 files stay readable (the block is simply
# absent) and converge to v2 on their next write; no bulk rewrite.
# v3 (issue #249): entries may carry an optional `facts` block —
# analyst corrections to LLM-extracted facts, keyed by fact name,
# applied as an overlay in `score aggregate` before the rule cascade
# runs. Same additive-convergent contract as v2.
CURATION_VERSION = 3

# Module constants baked at import — tests monkeypatch these
# (the `_OUT_DIR` convention from CLAUDE.md operational gotchas).
_CURATION_DIR = project_root() / "data" / "curation"
_OUT_DIR = out_dir()

# {header: key} for the analyst-input bands, derived from the spec so
# the ingest reader and the workbook renderer can never disagree about
# which columns are analyst-owned (single source of truth).
DETAILS_BAND_BY_HEADER: dict[str, str] = {
    c.header: c.key for c in DETAILS_COLUMNS if c.role == "analyst_input"
}
TRACKER_BAND_BY_HEADER: dict[str, str] = {
    c.header: c.key
    for c in COMMITMENT_TRACKER_COLUMNS
    if c.role == "analyst_input"
}

# Column keys coerced to float on ingest. Non-numeric input is captured
# verbatim as text but excluded from disagreement math (and counted in
# the ingest report as `non_numeric_override`).
_NUMERIC_KEYS = frozenset({"analyst_adjusted_override", "analyst_base_override"})

_IDENTITY_HEADERS = ("State", "Entity Name", "Data Element")
# Spine-lens Details carries this spine-only column; its presence is the
# lens discriminator for a workbook being ingested.
_SPINE_MARKER_HEADER = "AI: Documented"
_DETAILS_SHEET_TITLE = "Details"
_WORKBOOK_GENERATED_LABEL = "Workbook generated"
# Abort when more than this fraction of rows-with-input fail key
# validation — the column mapping is probably wrong, not the analyst.
_UNKNOWN_ABORT_FRACTION = 0.05
_HISTORY_CAP = 50

# Tolerance for float comparison between an analyst override and the
# engine score on either axis (0.5-step scale; anything past noise
# disagrees).
_OVERRIDE_EPSILON = 0.01

# The two override axes (issue #250): the headline adjusted score and
# the rule-cascade base tier, each with its own band column so the
# analyst can say precisely which layer they disagree with. Order here
# is render order — adjusted (the headline axis) sorts before base for
# a record contested on both. Each entry: (axis, curation key, engine
# reader over the score-sidecar record, why-text template).
_OVERRIDE_AXES: tuple[tuple[str, str, Callable[[dict], object], str], ...] = (
    (
        "adjusted",
        "analyst_adjusted_override",
        lambda s: s.get("adjusted_nachos_score"),
        "analyst override {override:g} vs engine adjusted {engine:g}",
    ),
    (
        "base",
        "analyst_base_override",
        lambda s: ((s.get("dimensions") or {}).get("nachos_score") or {}).get(
            "value"
        ),
        "analyst base override {override:g} vs engine base {engine:g}",
    ),
)


class IngestAbort(RuntimeError):
    """Raised when the ingest looks structurally wrong (bad sheet/columns
    or a mostly-unresolvable key set) — nothing is written."""


def curation_dir() -> Path:
    """The committed curation-sidecar directory (``data/curation/``)."""
    return _CURATION_DIR


def state_curation_path(state: str, base: Path | None = None) -> Path:
    """Return ``data/curation/{state}.json`` for the given state."""
    return (base or _CURATION_DIR) / f"{state.lower()}.json"


def _empty_payload(state: str) -> dict:
    return {
        "version": CURATION_VERSION,
        "state": state.upper(),
        "updated_at": None,
        "entries": {},
    }


def load_state_curation(state: str, base: Path | None = None) -> dict:
    """Full curation payload for a state (fresh skeleton when absent)."""
    path = state_curation_path(state, base)
    if not path.exists():
        return _empty_payload(state)
    return json.loads(path.read_text(encoding="utf-8"))


def curation_values_for(
    state: str, base: Path | None = None
) -> dict[str, dict[str, CellValue]]:
    """Flat ``{record_key: {column_key: value}}`` view for rendering.

    This is what ``analyst.run()`` threads into ``RowContext.curation``
    so the analyst-input band extractors re-apply stored entries on
    every regeneration.
    """
    payload = load_state_curation(state, base)
    out: dict[str, dict[str, CellValue]] = {}
    for record_key, entry in (payload.get("entries") or {}).items():
        values = {
            col_key: captured.get("value")
            for col_key, captured in (entry.get("values") or {}).items()
            if captured.get("value") not in (None, "")
        }
        if values:
            out[record_key] = values
    return out


# ---------------------------------------------------------------------------
# Ingest
# ---------------------------------------------------------------------------


@dataclass
class IngestReport:
    """What one ``review ingest`` run scanned, captured, and skipped."""

    workbook: str
    lens: str
    dry_run: bool
    author: str | None = None
    states: list[str] = field(default_factory=list)
    rows_scanned: int = 0
    rows_with_input: int = 0
    captured_per_column: dict[str, int] = field(default_factory=dict)
    replaced: int = 0
    unchanged: int = 0
    non_numeric_override: int = 0
    # (sheet row number, attempted record key) — listed so nothing
    # vanishes silently.
    unknown_rows: list[tuple[int, str]] = field(default_factory=list)
    duplicate_rows: list[tuple[int, str]] = field(default_factory=list)
    stale_warning: str | None = None
    written_paths: list[str] = field(default_factory=list)

    def render_text(self) -> str:
        lines = [
            f"review ingest — {self.workbook} "
            f"({self.lens} lens{', DRY RUN' if self.dry_run else ''})",
            f"  states: {', '.join(self.states) or '(none)'}",
            f"  rows scanned: {self.rows_scanned}; "
            f"rows with analyst input: {self.rows_with_input}",
        ]
        if self.captured_per_column:
            lines.append("  captured per column:")
            for col, n in sorted(self.captured_per_column.items()):
                lines.append(f"    {col}: {n}")
        else:
            lines.append("  captured per column: (nothing new)")
        lines.append(
            f"  new/updated values: "
            f"{sum(self.captured_per_column.values())} "
            f"(replaced {self.replaced} prior); unchanged: {self.unchanged}"
        )
        if self.non_numeric_override:
            lines.append(
                f"  WARNING: {self.non_numeric_override} override cell(s) "
                "were not numeric — captured as text, excluded from "
                "disagreement math"
            )
        if self.unknown_rows:
            lines.append(
                f"  WARNING: {len(self.unknown_rows)} row(s) with input did "
                "not match any known (state, entity, element) — skipped "
                "(identity cell edited, or a scratch row?):"
            )
            for row_idx, key in self.unknown_rows[:20]:
                lines.append(f"    sheet row {row_idx}: {key}")
            if len(self.unknown_rows) > 20:
                lines.append(f"    … +{len(self.unknown_rows) - 20} more")
        if self.duplicate_rows:
            lines.append(
                f"  WARNING: {len(self.duplicate_rows)} duplicate key row(s) "
                "— first occurrence kept:"
            )
            for row_idx, key in self.duplicate_rows[:10]:
                lines.append(f"    sheet row {row_idx}: {key}")
        if self.stale_warning:
            lines.append(f"  WARNING: {self.stale_warning}")
        if self.dry_run:
            lines.append("  dry run — nothing written")
        elif self.written_paths:
            lines.append("  wrote:")
            for p in self.written_paths:
                lines.append(f"    {p}")
        else:
            lines.append("  nothing to write")
        lines.append(
            "  note: blank cells never clear stored values (no delete "
            "mechanism in v1 — edit data/curation/*.json directly if you "
            "must remove an entry)"
        )
        return "\n".join(lines)


def _norm_header(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip())


def _norm_cell(value: object) -> CellValue:
    """JSON-serializable cell value; blank-ish → None."""
    if value is None:
        return None
    if isinstance(value, str):
        s = value.strip()
        return s or None
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, (int, float)):
        return value
    if isinstance(value, datetime):
        return value.isoformat(timespec="seconds")
    return str(value)


def _maybe_float(value: CellValue) -> float | None:
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        try:
            return float(value.strip())
        except ValueError:
            return None
    return None


def _now_iso() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def _valid_keys_for(state: str, lens: str, out_base: Path) -> set[str]:
    """Record keys the elements artifact knows for (state, lens)."""
    path = out_base / f"{state.lower()}_elements_{lens}.json"
    if not path.exists():
        raise IngestAbort(
            f"cannot validate {state} rows: {path} not found — run the "
            f"ingest pipeline first (keys are validated against the "
            f"elements artifact so edited identity cells can't corrupt "
            f"the curation sidecar)"
        )
    data = json.loads(path.read_text(encoding="utf-8"))
    st = state.upper()
    return {
        f"{st}|{r.get('entity')}|{r.get('element_name')}"
        for r in data.get("elements", [])
    }


def _workbook_generated_stamp(wb) -> str | None:
    """Best-effort read of the Update Log 'Workbook generated' cell."""
    for title in wb.sheetnames:
        if not title.startswith("Update Log"):
            continue
        ws = wb[title]
        for row in ws.iter_rows(min_row=1, max_row=30, max_col=2,
                                values_only=True):
            if row and _norm_header(row[0]) == _WORKBOOK_GENERATED_LABEL:
                return _norm_header(row[1]) if len(row) > 1 else None
    return None


def _stale_warning_for(
    states: list[str], lens: str, generated: str | None, out_base: Path
) -> str | None:
    if not generated:
        return None
    try:
        generated_at = datetime.fromisoformat(generated)
    except ValueError:
        return None
    for st in states:
        sidecar = out_base / f"{st.lower()}_scores_{lens}.json"
        if not sidecar.exists():
            continue
        sidecar_at = datetime.fromtimestamp(
            sidecar.stat().st_mtime, tz=timezone.utc
        )
        if generated_at < sidecar_at:
            return (
                f"workbook predates the current {st} scores sidecar "
                f"(workbook {generated_at.isoformat(timespec='seconds')} < "
                f"sidecar {sidecar_at.isoformat(timespec='seconds')}) — "
                "identity keys still join; Row #s may have shifted"
            )
    return None


_TRACKER_SHEET_TITLE = "Commitment Tracker"


def _band_sheets(wb) -> list[tuple[str, dict[str, str]]]:
    """(sheet title, {header: column key}) for every ingestable sheet:
    the Details analyst band + the Commitment Tracker's commitment
    columns (both keyed by the same record key into the same sidecar)."""
    sheets: list[tuple[str, dict[str, str]]] = []
    if _DETAILS_SHEET_TITLE in wb.sheetnames:
        sheets.append((_DETAILS_SHEET_TITLE, DETAILS_BAND_BY_HEADER))
    if _TRACKER_SHEET_TITLE in wb.sheetnames:
        sheets.append((_TRACKER_SHEET_TITLE, TRACKER_BAND_BY_HEADER))
    return sheets


def ingest_workbook(
    path: Path,
    *,
    author: str | None = None,
    dry_run: bool = False,
    out_base: Path | None = None,
    curation_base: Path | None = None,
) -> IngestReport:
    """Read the analyst-input band back into per-state curation sidecars.

    Raises ``IngestAbort`` on structural problems (missing Details sheet,
    missing identity columns, or >5% of rows-with-input failing key
    validation — a wrong column mapping, not analyst noise). Per-row
    problems (edited identity cells, scratch rows, duplicates) warn,
    skip, and are listed in the returned report.
    """
    out_base = out_base or _OUT_DIR
    curation_base = curation_base or _CURATION_DIR
    if not path.exists():
        raise FileNotFoundError(f"workbook not found: {path}")

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = _band_sheets(wb)
        if not sheets:
            raise IngestAbort(
                f"no '{_DETAILS_SHEET_TITLE}' sheet in {path.name} — "
                f"available: {wb.sheetnames}"
            )
        generated = _workbook_generated_stamp(wb)

        # ---- scan --------------------------------------------------------
        # captured[state][record_key][col_key] = value
        captured: dict[str, dict[str, dict[str, CellValue]]] = {}
        report = IngestReport(
            workbook=path.name, lens="source", dry_run=dry_run, author=author
        )
        seen_keys: set[str] = set()
        for sheet_title, band_by_header in sheets:
            ws = wb[sheet_title]
            header_row = next(
                ws.iter_rows(min_row=1, max_row=1, values_only=True), ()
            )
            headers = [_norm_header(h) for h in header_row]
            col_of = {h: i for i, h in enumerate(headers)}
            missing = [h for h in _IDENTITY_HEADERS if h not in col_of]
            if missing:
                raise IngestAbort(
                    f"sheet {sheet_title!r} is missing identity column(s) "
                    f"{missing} — cannot key rows"
                )
            if sheet_title == _DETAILS_SHEET_TITLE:
                report.lens = (
                    "spine" if _SPINE_MARKER_HEADER in col_of else "source"
                )
            band_cols = [
                (h, key, col_of[h])
                for h, key in band_by_header.items()
                if h in col_of
            ]
            state_i = col_of["State"]
            entity_i = col_of["Entity Name"]
            element_i = col_of["Data Element"]
            for row_idx, raw in enumerate(
                ws.iter_rows(min_row=2, values_only=True), start=2
            ):
                if raw is None or all(v in (None, "") for v in raw):
                    continue
                report.rows_scanned += 1
                state = _norm_header(
                    raw[state_i] if state_i < len(raw) else None
                ).upper()
                entity = _norm_header(
                    raw[entity_i] if entity_i < len(raw) else None
                )
                element = _norm_header(
                    raw[element_i] if element_i < len(raw) else None
                )
                values: dict[str, CellValue] = {}
                for _h, key, i in band_cols:
                    v = _norm_cell(raw[i] if i < len(raw) else None)
                    if v is None:
                        continue
                    if key in _NUMERIC_KEYS:
                        num = _maybe_float(v)
                        if num is not None:
                            v = num
                        else:
                            report.non_numeric_override += 1
                    values[key] = v
                if not values:
                    continue
                report.rows_with_input += 1
                record_key = f"{state}|{entity}|{element}"
                if not state or not entity or not element:
                    report.unknown_rows.append((row_idx, record_key))
                    continue
                dedupe_key = f"{sheet_title}::{record_key}"
                if dedupe_key in seen_keys:
                    report.duplicate_rows.append((row_idx, record_key))
                    continue
                seen_keys.add(dedupe_key)
                captured.setdefault(state, {}).setdefault(record_key, {}).update(
                    values
                )
    finally:
        wb.close()

    # ---- validate keys against the elements artifacts ---------------------
    validated: dict[str, dict[str, dict[str, CellValue]]] = {}
    for state, by_key in sorted(captured.items()):
        valid = _valid_keys_for(state, report.lens, out_base)
        for record_key, values in by_key.items():
            if record_key in valid:
                validated.setdefault(state, {})[record_key] = values
            else:
                report.unknown_rows.append((-1, record_key))
    # Deduplicate unknown listings (a scan-time unknown may also be here).
    report.unknown_rows = sorted(set(report.unknown_rows))
    if (
        report.rows_with_input
        and len(report.unknown_rows)
        > _UNKNOWN_ABORT_FRACTION * report.rows_with_input
    ):
        raise IngestAbort(
            f"{len(report.unknown_rows)} of {report.rows_with_input} rows "
            f"with analyst input failed key validation (> "
            f"{_UNKNOWN_ABORT_FRACTION:.0%}) — the column mapping is "
            "probably wrong; nothing written. First few: "
            + "; ".join(k for _r, k in report.unknown_rows[:5])
        )

    report.states = sorted(validated)
    report.stale_warning = _stale_warning_for(
        report.states, report.lens, generated, out_base
    )

    # ---- merge into per-state sidecars ------------------------------------
    ingested_at = _now_iso()
    for state in report.states:
        payload = load_state_curation(state, curation_base)
        entries = payload.setdefault("entries", {})
        dirty = False
        for record_key, values in sorted(validated[state].items()):
            entry = entries.setdefault(
                record_key, {"values": {}, "history": []}
            )
            for col_key, value in values.items():
                existing = entry["values"].get(col_key)
                if existing is not None and existing.get("value") == value:
                    report.unchanged += 1
                    continue
                if existing is not None:
                    entry["history"].append(
                        {
                            "column": col_key,
                            "value": existing.get("value"),
                            "author": existing.get("author"),
                            "replaced_at": ingested_at,
                            "source_workbook": existing.get(
                                "source_workbook"
                            ),
                        }
                    )
                    entry["history"] = entry["history"][-_HISTORY_CAP:]
                    report.replaced += 1
                entry["values"][col_key] = {
                    "value": value,
                    "author": author,
                    "ingested_at": ingested_at,
                    "source_workbook": path.name,
                    "workbook_generated": generated,
                    "lens": report.lens,
                }
                report.captured_per_column[col_key] = (
                    report.captured_per_column.get(col_key, 0) + 1
                )
                dirty = True
        if not dirty:
            continue
        payload["updated_at"] = ingested_at
        if not dry_run:
            out_path = state_curation_path(state, curation_base)
            out_path.parent.mkdir(parents=True, exist_ok=True)
            out_path.write_text(
                json.dumps(payload, indent=2, sort_keys=True) + "\n",
                encoding="utf-8",
            )
            report.written_paths.append(str(out_path))
            logger.info("review ingest: wrote %s", out_path)
    return report


# ---------------------------------------------------------------------------
# Override disagreement (§8.4 — disagreement is a review signal)
# ---------------------------------------------------------------------------


def override_disagreements_for(
    state: str,
    scores: dict[str, dict],
    lens: str,
    base: Path | None = None,
) -> list[dict]:
    """Same-lens override-vs-engine disagreements for one state.

    Loads the state's curation sidecar and keeps only override values
    captured on the SAME lens as the workbook being rendered (the band's
    workflow columns re-apply on both lenses — deliberately — but a
    score override is a judgment about one lens's engine score). Both
    override axes (adjusted + base, issue #250) are gated independently
    — each captured value carries its own lens stamp. The lens gate
    lives here because the flat ``curation_values_for`` view drops
    provenance.
    """
    payload = load_state_curation(state, base)
    flat: dict[str, dict[str, CellValue]] = {}
    for record_key, entry in (payload.get("entries") or {}).items():
        values = entry.get("values") or {}
        for _axis, key, _engine_of, _why in _OVERRIDE_AXES:
            cap = values.get(key)
            if not cap:
                continue
            if cap.get("lens") not in (None, lens):
                continue
            flat.setdefault(record_key, {})[key] = cap.get("value")
    return override_disagreements(flat, scores, lens)


def override_disagreements(
    curation: dict[str, dict[str, CellValue]],
    scores: dict[str, dict],
    lens: str,
) -> list[dict]:
    """Rows where an analyst override disagrees with the engine.

    Pure function over a flat curation view + the score sidecar for one
    state (lens gating happens in ``override_disagreements_for``, which
    reads provenance from the sidecar payload). Two axes (issue #250):
    the adjusted override compares against the headline
    ``adjusted_nachos_score``; the base override compares against the
    rule-cascade tier at ``dimensions.nachos_score.value``. A record
    contested on both axes yields one disagreement per axis (the
    ``axis`` field says which) so each ``why`` sentence stays
    unambiguous. The engine score is NEVER modified — the disagreement
    itself is the signal, rendered at the top of the Review Queue sheet
    under the workbook-only ``OVERRIDE`` route label (deliberately not
    part of ``review_queue.ROUTES``, which drives the JSON artifacts).
    """
    out: list[dict] = []
    for record_key, values in curation.items():
        score = scores.get(record_key)
        if not score:
            continue
        for axis, key, engine_of, why in _OVERRIDE_AXES:
            override = _maybe_float(values.get(key))
            if override is None:
                continue
            engine = engine_of(score)
            if not isinstance(engine, (int, float)):
                continue
            if abs(override - float(engine)) <= _OVERRIDE_EPSILON:
                continue
            out.append(
                {
                    "record_key": record_key,
                    "entity": score.get("entity"),
                    "element_name": score.get("element_name"),
                    "axis": axis,
                    "override": override,
                    "engine_value": float(engine),
                    "why": why.format(override=override, engine=float(engine)),
                }
            )
    out.sort(key=lambda d: (d["record_key"], d["axis"]))
    return out


# ---------------------------------------------------------------------------
# Adjudication overlay (issue #248 Part B — curation schema v2)
# ---------------------------------------------------------------------------
#
# Team consensus on one row's ADJUSTED score, recorded as a distinct,
# provenance-carrying layer: the engine score is NEVER mutated, and
# adjudications never feed prompts, rules, or tests (the GT boundary
# applies to our own team's consensus exactly as it applies to external
# reviewer workbooks). Adjusted axis ONLY by design: the base tier is
# pure rule-cascade output, so consensus that a tier is wrong is by
# definition a rule problem — it routes to the version-bump path via
# the override-clustering diagnostic, not adjudication. The `axis`
# field is recorded explicitly for forward-compat; readers skip any
# other value. Adjudication is the pressure valve for genuinely
# idiosyncratic rows, not the default path.


class AdjudicationAbort(RuntimeError):
    """Raised when an adjudication cannot be recorded safely (unknown
    record key, or no engine score to stamp) — nothing is written."""


@dataclass
class AdjudicationReport:
    """What one ``review adjudicate`` run recorded."""

    state: str
    record_key: str
    lens: str
    value: float
    agreed_by: tuple[str, ...]
    engine_score_at_decision: float
    plan_version_at_decision: str
    replaced_prior: bool
    dry_run: bool
    written_path: str | None = None

    def render_text(self) -> str:
        lines = [
            f"review adjudicate — {self.record_key} "
            f"({self.lens} lens{', DRY RUN' if self.dry_run else ''})",
            f"  adjudicated value: {self.value:g} "
            f"(engine adjusted at decision: "
            f"{self.engine_score_at_decision:g}, "
            f"plan v{self.plan_version_at_decision})",
            f"  agreed by: {', '.join(self.agreed_by)}",
        ]
        if self.replaced_prior:
            lines.append(
                "  replaced a prior adjudication (moved to history)"
            )
        if self.dry_run:
            lines.append("  dry run — nothing written")
        elif self.written_path:
            lines.append(f"  wrote: {self.written_path}")
        lines.append(
            "  note: the engine score is never modified — the consensus "
            "renders as `Effective Score (adjudicated)` beside it, and "
            "goes stale if the engine or plan version moves"
        )
        return "\n".join(lines)


def adjudicate(
    state: str,
    entity: str,
    element: str,
    *,
    value: float,
    agreed_by: tuple[str, ...],
    rationale: str,
    lens: str = "source",
    allow_stale: bool = False,
    dry_run: bool = False,
    out_base: Path | None = None,
    curation_base: Path | None = None,
) -> AdjudicationReport:
    """Record team consensus for one row's adjusted score (schema v2).

    Stamps the CURRENT engine ``adjusted_nachos_score`` and
    ``SCORING_PLAN_VERSION`` into the block so staleness is detectable:
    if either moves later, the adjudication renders blank and the row
    re-enters the Review Queue as RE-ADJUDICATE. The engine score is
    read through the freshness-gated sidecar loader, so a drifted
    engine value can never be stamped silently (``allow_stale`` to
    override, mirroring the render path).
    """
    # Local import — loaders is render-layer, imported lazily like the
    # analyst sheet builders do with this module.
    from src.report.loaders import load_scores_sidecar

    out_base = out_base or _OUT_DIR
    st = state.upper()
    record_key = f"{st}|{entity}|{element}"
    try:
        valid = _valid_keys_for(st, lens, out_base)
    except IngestAbort as exc:
        raise AdjudicationAbort(str(exc)) from exc
    if record_key not in valid:
        siblings = sorted(
            k for k in valid if k.startswith(f"{st}|{entity}|")
        )[:5]
        hint = (
            "; known keys for this entity include: " + "; ".join(siblings)
            if siblings
            else ""
        )
        raise AdjudicationAbort(
            f"unknown record key {record_key!r} for the {lens} lens — "
            f"nothing written{hint}"
        )

    scores = load_scores_sidecar(st, lens, out_base, allow_stale=allow_stale)
    score = scores.get(record_key)
    engine = (score or {}).get("adjusted_nachos_score")
    if not isinstance(engine, (int, float)):
        raise AdjudicationAbort(
            f"cannot adjudicate an unscored record: {record_key} has no "
            f"numeric adjusted_nachos_score in the {lens} scores sidecar "
            "— engine_score_at_decision would be meaningless"
        )

    from src.score.aggregate import SCORING_PLAN_VERSION

    names = tuple(n.strip() for n in agreed_by if n and n.strip())
    if not names:
        raise AdjudicationAbort("at least one --agreed-by name is required")

    now = _now_iso()
    payload = load_state_curation(st, curation_base)
    entries = payload.setdefault("entries", {})
    entry = entries.setdefault(record_key, {"values": {}, "history": []})
    prior = entry.get("adjudication")
    if prior is not None:
        # Same newest-wins + history pattern as `values`: the entire
        # prior block is preserved under the `adjudication` column
        # discriminator (history stays one homogeneous capped list).
        entry.setdefault("history", []).append(
            {
                "column": "adjudication",
                "value": prior,
                "replaced_at": now,
            }
        )
        entry["history"] = entry["history"][-_HISTORY_CAP:]
    entry["adjudication"] = {
        "status": "adjudicated",
        "axis": "adjusted",
        "value": float(value),
        "lens": lens,
        "agreed_by": list(names),
        "decided_at": now,
        "rationale": rationale,
        "engine_score_at_decision": float(engine),
        "plan_version_at_decision": SCORING_PLAN_VERSION,
    }
    payload["version"] = CURATION_VERSION
    payload["updated_at"] = now

    report = AdjudicationReport(
        state=st,
        record_key=record_key,
        lens=lens,
        value=float(value),
        agreed_by=names,
        engine_score_at_decision=float(engine),
        plan_version_at_decision=SCORING_PLAN_VERSION,
        replaced_prior=prior is not None,
        dry_run=dry_run,
    )
    if not dry_run:
        out_path = state_curation_path(st, curation_base)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        out_path.write_text(
            json.dumps(payload, indent=2, sort_keys=True) + "\n",
            encoding="utf-8",
        )
        report.written_path = str(out_path)
        logger.info("review adjudicate: wrote %s", out_path)
    return report


def adjudications_for(
    state: str,
    scores: dict[str, dict],
    lens: str,
    base: Path | None = None,
) -> dict[str, dict]:
    """Resolved same-lens adjudications for one state, keyed by record.

    Each returned dict carries the block's fields verbatim plus
    computed render state: ``entity`` / ``element_name`` (from the
    score record, else parsed from the key), the CURRENT
    ``engine_value``, and ``fresh`` / ``stale_reason``. Fresh iff the
    record is scored with a numeric adjusted score, the engine hasn't
    moved beyond ``_OVERRIDE_EPSILON`` since the decision, and the plan
    version matches. A stale adjudication renders BLANK (never a stale
    consensus) and re-enters the queue as RE-ADJUDICATE. An
    adjudication that AGREES with the current engine is fresh and still
    renders — blank means "none or stale" only, and the freshness stamp
    re-arms if the engine later moves. Lens-gated like overrides
    (adjudication is a judgment about one lens's engine score); blocks
    with an axis other than ``adjusted`` are skipped (forward-compat).
    """
    from src.score.aggregate import SCORING_PLAN_VERSION

    payload = load_state_curation(state, base)
    out: dict[str, dict] = {}
    for record_key, entry in (payload.get("entries") or {}).items():
        block = entry.get("adjudication")
        if not block:
            continue
        if block.get("axis") != "adjusted" or block.get("lens") != lens:
            continue
        score = scores.get(record_key) or {}
        engine = score.get("adjusted_nachos_score")
        engine_value = (
            float(engine) if isinstance(engine, (int, float)) else None
        )
        at_decision = block.get("engine_score_at_decision")
        plan_at_decision = block.get("plan_version_at_decision")
        stale_reason: str | None = None
        if engine_value is None:
            stale_reason = "record no longer scored at the current plan"
        elif not isinstance(at_decision, (int, float)) or abs(
            engine_value - float(at_decision)
        ) > _OVERRIDE_EPSILON:
            stale_reason = (
                f"engine adjusted moved "
                f"{float(at_decision):g} → {engine_value:g}"
                if isinstance(at_decision, (int, float))
                else "no engine score recorded at decision time"
            )
        elif plan_at_decision != SCORING_PLAN_VERSION:
            stale_reason = (
                f"plan v{plan_at_decision} → v{SCORING_PLAN_VERSION}"
            )
        key_parts = record_key.split("|")
        out[record_key] = {
            **block,
            "entity": score.get("entity")
            or (key_parts[1] if len(key_parts) == 3 else None),
            "element_name": score.get("element_name")
            or (key_parts[2] if len(key_parts) == 3 else None),
            "engine_value": engine_value,
            "fresh": stale_reason is None,
            "stale_reason": stale_reason,
        }
    return out


# ---------------------------------------------------------------------------
# Fact-level curation (issue #249 — curation schema v3)
# ---------------------------------------------------------------------------
#
# The Case-1 branch of the override triage: when the engine's score is
# wrong because an EXTRACTED FACT is wrong, the honest lever is to
# correct the fact itself and let the unchanged rule cascade recompute
# — not to override the score downstream of the rules. Corrections are
# recorded here with full provenance and applied as an overlay in
# `score aggregate` (after fact loading, before rules), where the
# fact's sidecar provenance becomes `human_corrected`. The prompt
# cache is NEVER written: the cache stays the immutable record of what
# the model said; the overlay is the record of what the human said.
# v1 scope guard: LLM-extracted facts only — a wrong deterministic
# fact means the code is wrong (bug fix + DETERMINISTIC_VERSION bump,
# not a curation entry). GT boundary: corrections never seed prompts,
# rules, or tests; a correction PATTERN (many rows, one fact) is a
# prompt weakness and routes to the prompt-version-bump path.


class FactCorrectionAbort(RuntimeError):
    """Raised when a fact correction cannot be recorded safely (unknown
    key, unscored record, non-LLM fact, invalid value) — nothing is
    written."""


@dataclass
class FactCorrectionReport:
    """What one ``review correct-fact`` run recorded."""

    state: str
    record_key: str
    lens: str
    fact: str
    value: bool | int | str
    prior_value: bool | int | str | None
    prior_provenance: str
    author: str
    plan_version_at_correction: str
    replaced_prior: bool
    dry_run: bool
    written_path: str | None = None

    def render_text(self) -> str:
        lines = [
            f"review correct-fact — {self.record_key} "
            f"({self.lens} lens{', DRY RUN' if self.dry_run else ''})",
            f"  {self.fact}: {self.prior_value!r} ({self.prior_provenance})"
            f" → {self.value!r} (human_corrected)",
            f"  author: {self.author} "
            f"(plan v{self.plan_version_at_correction})",
        ]
        if self.replaced_prior:
            lines.append(
                "  replaced a prior correction for this fact "
                "(moved to history)"
            )
        if self.dry_run:
            lines.append("  dry run — nothing written")
        elif self.written_path:
            lines.append(f"  wrote: {self.written_path}")
        lines.append(
            "  note: the correction applies at aggregate time — run "
            f"`poc3 score aggregate --state {self.state.lower()} --lens "
            f"{self.lens}` (or `poc3 publish`) to recompute scores; the "
            "prompt cache is never touched"
        )
        return "\n".join(lines)


def correct_fact(
    state: str,
    entity: str,
    element: str,
    *,
    fact: str,
    value: str,
    rationale: str,
    author: str,
    lens: str = "source",
    allow_stale: bool = False,
    dry_run: bool = False,
    out_base: Path | None = None,
    curation_base: Path | None = None,
) -> FactCorrectionReport:
    """Record an analyst correction to one LLM-extracted fact (schema v3).

    Validates the record key against the elements artifact, reads the
    CURRENT extracted value through the freshness-gated sidecar loader
    (``allow_stale`` to override), and refuses anything the extraction
    schema would have refused from the model: unknown facts,
    deterministic facts (code-bug path, not curation), facts never
    extracted for this record/lens, and type-invalid values. Stamps
    ``prior_value`` / ``prior_provenance`` + the plan version at
    correction time as provenance (NOT staleness gates — the overlay
    always applies; re-aggregation is what moves the score).
    """
    from src.report.loaders import load_scores_sidecar
    from src.score import deterministic, extract
    from src.score.schema import parse_fact_value

    out_base = out_base or _OUT_DIR
    st = state.upper()
    record_key = f"{st}|{entity}|{element}"
    try:
        valid = _valid_keys_for(st, lens, out_base)
    except IngestAbort as exc:
        raise FactCorrectionAbort(str(exc)) from exc
    if record_key not in valid:
        siblings = sorted(
            k for k in valid if k.startswith(f"{st}|{entity}|")
        )[:5]
        hint = (
            "; known keys for this entity include: " + "; ".join(siblings)
            if siblings
            else ""
        )
        raise FactCorrectionAbort(
            f"unknown record key {record_key!r} for the {lens} lens — "
            f"nothing written{hint}"
        )

    if fact in deterministic.DETERMINISTIC_FACTS:
        raise FactCorrectionAbort(
            f"{fact} is a deterministic fact — a wrong value there means "
            "the computation is wrong. That is a code bug (fix it and "
            f"bump DETERMINISTIC_VERSION, currently "
            f"{deterministic.DETERMINISTIC_VERSION}), not a curation "
            "entry. Corrections cover LLM-extracted facts only."
        )
    if fact not in extract.SUPPORTED_FACTS:
        raise FactCorrectionAbort(
            f"unknown fact {fact!r} — correctable LLM facts are: "
            + ", ".join(extract.SUPPORTED_FACTS)
        )

    scores = load_scores_sidecar(st, lens, out_base, allow_stale=allow_stale)
    score = scores.get(record_key)
    if score is None:
        raise FactCorrectionAbort(
            f"cannot correct an unscored record: {record_key} is not in "
            f"the {lens} scores sidecar — prior_value would be meaningless"
        )
    prov = (score.get("fact_provenance") or {}).get(fact)
    if prov is None:
        raise FactCorrectionAbort(
            f"{fact} was not extracted for {record_key} on the {lens} "
            "lens — a correction here would invent a fact, not correct "
            "one"
        )
    reason = prov.get("downgrade_reason")
    if reason == "filtered_by_source":
        raise FactCorrectionAbort(
            f"{fact} does not apply to {record_key} (source-filtered: "
            "the fact only runs on extension rows) — nothing to correct"
        )
    if reason == "missing_from_artifact":
        hint = ""
        if lens == "spine" and fact == "extension_is_necessary":
            hint = (
                " (spine scoring borrows this fact from the SOURCE "
                "sidecar — correct the source-lens row instead)"
            )
        raise FactCorrectionAbort(
            f"{fact} was never extracted for {record_key} on the {lens} "
            f"lens{hint} — a correction here would invent a fact, not "
            "correct one"
        )

    try:
        corrected = parse_fact_value(fact, value)
    except ValueError as exc:
        raise FactCorrectionAbort(str(exc)) from exc

    payload = load_state_curation(st, curation_base)
    entries = payload.setdefault("entries", {})
    entry = entries.setdefault(record_key, {"values": {}, "history": []})
    facts = entry.setdefault("facts", {})
    prior_block = facts.get(fact)

    # The no-op guard compares against the EFFECTIVE value — the
    # standing correction when one exists (the overlay always applies
    # the newest block), else the extracted value. This is what makes
    # "correct back to the model's value" possible: the analyst's
    # retract-equivalent when the model was right after all.
    prior_value = prov.get("value")
    effective = (
        prior_block.get("value") if prior_block is not None else prior_value
    )
    if corrected == effective:
        raise FactCorrectionAbort(
            f"{fact} already reads {effective!r} for {record_key} on "
            f"the {lens} lens — nothing to correct"
        )
    if prov.get("provenance") == "human_corrected":
        # The sidecar was re-aggregated under a prior correction; the
        # value being replaced is the human's, not the model's.
        prior_provenance = "human_corrected"
    elif prov.get("downgraded"):
        prior_provenance = f"llm_downgraded:{reason}"
    else:
        prior_provenance = "llm"

    author = author.strip()
    if not author:
        raise FactCorrectionAbort("--author is required")
    rationale = rationale.strip()
    if not rationale:
        raise FactCorrectionAbort("--rationale is required")

    from src.score.aggregate import SCORING_PLAN_VERSION

    now = _now_iso()
    if prior_block is not None:
        # Same newest-wins + history pattern as `values` and
        # `adjudication`: the entire prior block is preserved under a
        # `fact:{name}` column discriminator.
        entry.setdefault("history", []).append(
            {
                "column": f"fact:{fact}",
                "value": prior_block,
                "replaced_at": now,
            }
        )
        entry["history"] = entry["history"][-_HISTORY_CAP:]
    facts[fact] = {
        "value": corrected,
        "lens": lens,
        "author": author,
        "corrected_at": now,
        "rationale": rationale,
        "prior_value": prior_value,
        "prior_provenance": prior_provenance,
        "plan_version_at_correction": SCORING_PLAN_VERSION,
    }
    payload["version"] = CURATION_VERSION
    payload["updated_at"] = now

    report = FactCorrectionReport(
        state=st,
        record_key=record_key,
        lens=lens,
        fact=fact,
        value=corrected,
        prior_value=prior_value,
        prior_provenance=prior_provenance,
        author=author,
        plan_version_at_correction=SCORING_PLAN_VERSION,
        replaced_prior=prior_block is not None,
        dry_run=dry_run,
    )
    if not dry_run:
        out_path = state_curation_path(st, curation_base)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        out_path.write_text(
            json.dumps(payload, indent=2, sort_keys=True) + "\n",
            encoding="utf-8",
        )
        report.written_path = str(out_path)
        logger.info("review correct-fact: wrote %s", out_path)
    return report


def fact_corrections_for(
    state: str,
    lens: str,
    base: Path | None = None,
) -> dict[str, dict[str, dict]]:
    """Same-lens fact corrections for one state, keyed by record.

    Returns ``{record_key: {fact_name: correction_block}}`` filtered to
    blocks whose stored ``lens`` stamp matches — a correction is a
    statement about the lens it was captured on (the spine lens picks
    up a corrected source ``extension_is_necessary`` anyway, through
    the existing sidecar borrow). This is what the aggregate overlay
    consumes; blocks are returned verbatim (the overlay re-validates
    values before applying, so a hand-edited sidecar degrades loudly
    rather than crashing the run).
    """
    payload = load_state_curation(state, base)
    out: dict[str, dict[str, dict]] = {}
    for record_key, entry in (payload.get("entries") or {}).items():
        blocks = {
            fact: block
            for fact, block in (entry.get("facts") or {}).items()
            if block.get("lens") == lens
        }
        if blocks:
            out[record_key] = blocks
    return out


# ---------------------------------------------------------------------------
# CLI entry (plain function — cli.py wraps it; never @click.command)
# ---------------------------------------------------------------------------


def run(
    workbook: str | Path,
    *,
    author: str | None = None,
    dry_run: bool = False,
) -> IngestReport:
    """Ingest analyst edits from one workbook. Returns the report."""
    return ingest_workbook(Path(workbook), author=author, dry_run=dry_run)

"""Tests for the Option C analyst round-trip (`mc review ingest` +
curation sidecar re-apply). Issue #186 sequence 3."""

from __future__ import annotations

import inspect
import json
from datetime import datetime, timedelta, timezone
from pathlib import Path

import click
import openpyxl
import pytest
from openpyxl import Workbook

from src.report import curation
from src.report.curation import (
    DETAILS_BAND_BY_HEADER,
    IngestAbort,
    curation_values_for,
    ingest_workbook,
    override_disagreements,
    state_curation_path,
)


class TestCliWiring:
    def test_run_is_plain_function_not_click_command(self):
        assert not isinstance(curation.run, click.Command)
        assert inspect.isfunction(curation.run)

    def test_cli_command_calls_module_run(self, monkeypatch, tmp_path):
        from click.testing import CliRunner

        from src.cli import cli

        calls: dict = {}

        class _FakeReport:
            def render_text(self):
                return "ok"

        def fake_run(workbook, *, author=None, dry_run=False):
            calls["workbook"] = Path(workbook)
            calls["author"] = author
            calls["dry_run"] = dry_run
            return _FakeReport()

        monkeypatch.setattr(curation, "run", fake_run)
        wb_path = tmp_path / "az_analyst.xlsx"
        Workbook().save(wb_path)
        result = CliRunner().invoke(
            cli,
            ["review", "ingest", str(wb_path), "--author", "maria",
             "--dry-run"],
        )
        assert result.exit_code == 0, result.output
        assert calls == {
            "workbook": wb_path, "author": "maria", "dry_run": True,
        }
        assert "ok" in result.output

    def test_adjudicate_is_plain_function_not_click_command(self):
        assert not isinstance(curation.adjudicate, click.Command)
        assert inspect.isfunction(curation.adjudicate)

    def test_cli_adjudicate_calls_module_function(self, monkeypatch):
        from click.testing import CliRunner

        from src.cli import cli

        calls: dict = {}

        class _FakeReport:
            def render_text(self):
                return "recorded"

        def fake_adjudicate(state, entity, element, **kw):
            calls["key"] = (state, entity, element)
            calls.update(kw)
            return _FakeReport()

        monkeypatch.setattr(curation, "adjudicate", fake_adjudicate)
        result = CliRunner().invoke(
            cli,
            ["review", "adjudicate", "--state", "az",
             "--entity", "Calendar", "--element", "calendarCode",
             "--value", "2.0", "--agreed-by", "Doug",
             "--agreed-by", "Maria", "--rationale", "consensus 07-13",
             "--dry-run"],
        )
        assert result.exit_code == 0, result.output
        assert calls == {
            "key": ("az", "Calendar", "calendarCode"),
            "value": 2.0,
            "agreed_by": ("Doug", "Maria"),
            "rationale": "consensus 07-13",
            "lens": "source",
            "allow_stale": False,
            "dry_run": True,
        }
        assert "recorded" in result.output

    def test_correct_fact_is_plain_function_not_click_command(self):
        assert not isinstance(curation.correct_fact, click.Command)
        assert inspect.isfunction(curation.correct_fact)

    def test_cli_correct_fact_calls_module_function(self, monkeypatch):
        from click.testing import CliRunner

        from src.cli import cli

        calls: dict = {}

        class _FakeReport:
            def render_text(self):
                return "corrected"

        def fake_correct_fact(state, entity, element, **kw):
            calls["key"] = (state, entity, element)
            calls.update(kw)
            return _FakeReport()

        monkeypatch.setattr(curation, "correct_fact", fake_correct_fact)
        result = CliRunner().invoke(
            cli,
            ["review", "correct-fact", "--state", "tx",
             "--entity", "Assessment", "--element", "AcademicSubject",
             "--fact", "has_conditional_logic", "--value", "false",
             "--author", "Chris Moffatt",
             "--rationale", "calendar, not conditionality",
             "--dry-run"],
        )
        assert result.exit_code == 0, result.output
        assert calls == {
            "key": ("tx", "Assessment", "AcademicSubject"),
            "fact": "has_conditional_logic",
            "value": "false",
            "author": "Chris Moffatt",
            "rationale": "calendar, not conditionality",
            "lens": "source",
            "allow_stale": False,
            "dry_run": True,
        }
        assert "corrected" in result.output


# ---------------------------------------------------------------------------
# Fixture helpers — synthetic workbook + elements artifacts in tmp_path
# ---------------------------------------------------------------------------

_BAND_HEADERS = tuple(DETAILS_BAND_BY_HEADER)  # spec order


def _write_elements(out: Path, state: str, pairs: list[tuple[str, str]],
                    lens: str = "source") -> None:
    out.mkdir(parents=True, exist_ok=True)
    (out / f"{state.lower()}_elements_{lens}.json").write_text(
        json.dumps(
            {
                "state": state,
                "elements": [
                    {"entity": e, "element_name": n} for e, n in pairs
                ],
            }
        ),
        encoding="utf-8",
    )


def _make_workbook(
    path: Path,
    rows: list[dict],
    *,
    spine: bool = False,
    generated: str | None = None,
) -> None:
    """Details-sheet workbook in the generated-analyst shape (subset)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Details"
    headers = ["Row #", "State", "Entity Name", "Data Element",
               "Adjusted NACHOS Score"]
    if spine:
        headers.append("AI: Documented")
    headers.extend(_BAND_HEADERS)
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h) for h in headers])
    if generated is not None:
        ul = wb.create_sheet("Update Log")
        ul.cell(row=1, column=1, value="Update Log — TEST")
        ul.cell(row=3, column=1, value="Workbook generated")
        ul.cell(row=3, column=2, value=generated)
    wb.save(path)


@pytest.fixture
def env(tmp_path):
    """(out_dir, curation_dir, workbook_path) with a 2-row AZ universe."""
    out = tmp_path / "out"
    cur = tmp_path / "curation"
    _write_elements(out, "AZ", [("Calendar", "calendarCode"),
                                ("School", "schoolId")])
    return out, cur, tmp_path / "az_analyst.xlsx"


def _ingest(path, out, cur, **kw):
    return ingest_workbook(path, out_base=out, curation_base=cur, **kw)


# ---------------------------------------------------------------------------
# Capture + merge semantics
# ---------------------------------------------------------------------------


class TestIngestCapture:
    def test_captures_values_with_provenance(self, env):
        out, cur, wb = env
        _make_workbook(wb, [
            {"State": "AZ", "Entity Name": "Calendar",
             "Data Element": "calendarCode",
             "Reviewed?": "Yes", "Ed-Fi Comments": "looks right"},
            {"State": "AZ", "Entity Name": "School",
             "Data Element": "schoolId"},  # no analyst input
        ])
        report = _ingest(wb, out, cur, author="maria")
        assert report.states == ["AZ"]
        assert report.rows_scanned == 2
        assert report.rows_with_input == 1
        assert report.captured_per_column == {
            "reviewed": 1, "edfi_comments": 1,
        }
        payload = json.loads(
            state_curation_path("AZ", cur).read_text(encoding="utf-8")
        )
        entry = payload["entries"]["AZ|Calendar|calendarCode"]
        cap = entry["values"]["reviewed"]
        assert cap["value"] == "Yes"
        assert cap["author"] == "maria"
        assert cap["lens"] == "source"
        assert cap["source_workbook"] == "az_analyst.xlsx"
        assert cap["ingested_at"]
        assert entry["history"] == []
        # Untouched row never enters the sidecar.
        assert "AZ|School|schoolId" not in payload["entries"]

    def test_reingest_identical_values_is_a_noop(self, env):
        out, cur, wb = env
        _make_workbook(wb, [
            {"State": "AZ", "Entity Name": "Calendar",
             "Data Element": "calendarCode", "Reviewed?": "Yes"},
        ])
        _ingest(wb, out, cur)
        before = state_curation_path("AZ", cur).read_text(encoding="utf-8")
        report = _ingest(wb, out, cur)
        assert report.unchanged == 1
        assert report.captured_per_column == {}
        assert report.written_paths == []
        assert state_curation_path("AZ", cur).read_text(
            encoding="utf-8") == before

    def test_newest_wins_and_history_keeps_prior(self, env):
        out, cur, wb = env
        _make_workbook(wb, [
            {"State": "AZ", "Entity Name": "Calendar",
             "Data Element": "calendarCode", "Ed-Fi Comments": "first pass"},
        ])
        _ingest(wb, out, cur, author="maria")
        _make_workbook(wb, [
            {"State": "AZ", "Entity Name": "Calendar",
             "Data Element": "calendarCode",
             "Ed-Fi Comments": "second pass"},
        ])
        report = _ingest(wb, out, cur, author="doug")
        assert report.replaced == 1
        payload = json.loads(
            state_curation_path("AZ", cur).read_text(encoding="utf-8")
        )
        entry = payload["entries"]["AZ|Calendar|calendarCode"]
        assert entry["values"]["edfi_comments"]["value"] == "second pass"
        assert entry["values"]["edfi_comments"]["author"] == "doug"
        [hist] = entry["history"]
        assert hist["column"] == "edfi_comments"
        assert hist["value"] == "first pass"
        assert hist["author"] == "maria"

    def test_blank_cell_never_clears_stored_value(self, env):
        out, cur, wb = env
        _make_workbook(wb, [
            {"State": "AZ", "Entity Name": "Calendar",
             "Data Element": "calendarCode",
             "Ed-Fi Comments": "keep me", "Reviewed?": "Yes"},
        ])
        _ingest(wb, out, cur)
        # Analyst re-sends the workbook with the comment cell blanked
        # (filtered view, accidental delete — doesn't matter).
        _make_workbook(wb, [
            {"State": "AZ", "Entity Name": "Calendar",
             "Data Element": "calendarCode", "Reviewed?": "Yes"},
        ])
        _ingest(wb, out, cur)
        values = curation_values_for("AZ", cur)["AZ|Calendar|calendarCode"]
        assert values["edfi_comments"] == "keep me"
        assert values["reviewed"] == "Yes"

    def test_dry_run_writes_nothing(self, env):
        out, cur, wb = env
        _make_workbook(wb, [
            {"State": "AZ", "Entity Name": "Calendar",
             "Data Element": "calendarCode", "Reviewed?": "Yes"},
        ])
        report = _ingest(wb, out, cur, dry_run=True)
        assert report.dry_run is True
        assert report.captured_per_column == {"reviewed": 1}
        assert not cur.exists()
        assert "DRY RUN" in report.render_text()

    def test_numeric_override_coercion(self, env):
        out, cur, wb = env
        _make_workbook(wb, [
            {"State": "AZ", "Entity Name": "Calendar",
             "Data Element": "calendarCode",
             "Analyst Adjusted Score (override)": "1.5"},
            {"State": "AZ", "Entity Name": "School",
             "Data Element": "schoolId",
             "Analyst Adjusted Score (override)": "1.0 - see comment"},
        ])
        report = _ingest(wb, out, cur)
        assert report.non_numeric_override == 1
        values = curation_values_for("AZ", cur)
        assert values["AZ|Calendar|calendarCode"][
            "analyst_adjusted_override"] == 1.5
        # Non-numeric survives as text (excluded from disagreement math).
        assert values["AZ|School|schoolId"][
            "analyst_adjusted_override"] == "1.0 - see comment"

    def test_numeric_base_override_coercion(self, env):
        """The base-axis override (issue #250) joins `_NUMERIC_KEYS` —
        same coercion + non-numeric accounting as the adjusted axis."""
        out, cur, wb = env
        _make_workbook(wb, [
            {"State": "AZ", "Entity Name": "Calendar",
             "Data Element": "calendarCode",
             "Analyst Base Score (override)": "1.5"},
            {"State": "AZ", "Entity Name": "School",
             "Data Element": "schoolId",
             "Analyst Base Score (override)": "tier 2, surely"},
        ])
        report = _ingest(wb, out, cur)
        assert report.non_numeric_override == 1
        values = curation_values_for("AZ", cur)
        # Half-tier judgments are accepted (flag-don't-reject: no
        # integer enforcement on the base axis).
        assert values["AZ|Calendar|calendarCode"][
            "analyst_base_override"] == 1.5
        assert values["AZ|School|schoolId"][
            "analyst_base_override"] == "tier 2, surely"


class TestIngestValidation:
    def test_unknown_key_warns_and_skips_below_threshold(self, tmp_path):
        out = tmp_path / "out"
        cur = tmp_path / "curation"
        pairs = [("Entity", f"el{i}") for i in range(21)]
        _write_elements(out, "AZ", pairs)
        rows = [
            {"State": "AZ", "Entity Name": "Entity",
             "Data Element": f"el{i}", "Reviewed?": "Yes"}
            for i in range(21)
        ]
        # One row whose identity cell the analyst edited.
        rows.append({"State": "AZ", "Entity Name": "EntityRenamed",
                     "Data Element": "el0", "Reviewed?": "Yes"})
        wb = tmp_path / "az_analyst.xlsx"
        _make_workbook(wb, rows)
        report = _ingest(wb, out, cur)
        assert [k for _r, k in report.unknown_rows] == [
            "AZ|EntityRenamed|el0"
        ]
        assert "AZ|EntityRenamed|el0" not in curation_values_for("AZ", cur)
        assert "did not match" in report.render_text()

    def test_mostly_unknown_aborts_without_writing(self, env):
        out, cur, wb = env
        _make_workbook(wb, [
            {"State": "AZ", "Entity Name": "Nope",
             "Data Element": "wrong", "Reviewed?": "Yes"},
        ])
        with pytest.raises(IngestAbort, match="column mapping"):
            _ingest(wb, out, cur)
        assert not cur.exists()

    def test_missing_elements_artifact_aborts(self, tmp_path):
        out = tmp_path / "out"
        out.mkdir()
        wb = tmp_path / "az_analyst.xlsx"
        _make_workbook(wb, [
            {"State": "AZ", "Entity Name": "Calendar",
             "Data Element": "calendarCode", "Reviewed?": "Yes"},
        ])
        with pytest.raises(IngestAbort, match="elements"):
            _ingest(wb, out, tmp_path / "curation")

    def test_duplicate_key_keeps_first(self, env):
        out, cur, wb = env
        _make_workbook(wb, [
            {"State": "AZ", "Entity Name": "Calendar",
             "Data Element": "calendarCode", "Ed-Fi Comments": "first"},
            {"State": "AZ", "Entity Name": "Calendar",
             "Data Element": "calendarCode", "Ed-Fi Comments": "second"},
        ])
        report = _ingest(wb, out, cur)
        assert len(report.duplicate_rows) == 1
        assert curation_values_for("AZ", cur)[
            "AZ|Calendar|calendarCode"]["edfi_comments"] == "first"

    def test_missing_details_sheet_aborts(self, env):
        out, cur, _ = env
        path = out.parent / "not_analyst.xlsx"
        wb = Workbook()
        wb.active.title = "Nothing Here"
        wb.save(path)
        with pytest.raises(IngestAbort, match="Details"):
            _ingest(path, out, cur)

    def test_missing_identity_column_aborts(self, env):
        out, cur, _ = env
        path = out.parent / "broken.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Details"
        ws.append(["State", "Entity Name", "Reviewed?"])  # no Data Element
        ws.append(["AZ", "Calendar", "Yes"])
        wb.save(path)
        with pytest.raises(IngestAbort, match="identity"):
            _ingest(path, out, cur)


class TestIngestLensAndScope:
    def test_spine_marker_detects_lens(self, tmp_path):
        out = tmp_path / "out"
        cur = tmp_path / "curation"
        # Keys valid on the SPINE artifact only — proves validation ran
        # against the detected lens.
        _write_elements(out, "AZ", [("Calendar", "calendarCode")],
                        lens="spine")
        wb = tmp_path / "az_analyst_spine.xlsx"
        _make_workbook(wb, [
            {"State": "AZ", "Entity Name": "Calendar",
             "Data Element": "calendarCode", "Reviewed?": "Yes",
             "AI: Documented": "Yes"},
        ], spine=True)
        report = _ingest(wb, out, cur)
        assert report.lens == "spine"
        cap = json.loads(
            state_curation_path("AZ", cur).read_text(encoding="utf-8")
        )["entries"]["AZ|Calendar|calendarCode"]["values"]["reviewed"]
        assert cap["lens"] == "spine"

    def test_combined_workbook_splits_per_state(self, tmp_path):
        out = tmp_path / "out"
        cur = tmp_path / "curation"
        _write_elements(out, "AZ", [("Calendar", "calendarCode")])
        _write_elements(out, "MN", [("School", "schoolId")])
        wb = tmp_path / "coverage_analyst.xlsx"
        _make_workbook(wb, [
            {"State": "AZ", "Entity Name": "Calendar",
             "Data Element": "calendarCode", "Reviewed?": "Yes"},
            {"State": "MN", "Entity Name": "School",
             "Data Element": "schoolId", "DS Next Steps": "call MDE"},
        ])
        report = _ingest(wb, out, cur)
        assert report.states == ["AZ", "MN"]
        assert curation_values_for("AZ", cur)[
            "AZ|Calendar|calendarCode"]["reviewed"] == "Yes"
        assert curation_values_for("MN", cur)[
            "MN|School|schoolId"]["ds_next_steps"] == "call MDE"

    def test_stale_workbook_warning(self, env):
        out, cur, wb = env
        old = (datetime.now(timezone.utc) - timedelta(days=7)).isoformat(
            timespec="seconds"
        )
        _make_workbook(wb, [
            {"State": "AZ", "Entity Name": "Calendar",
             "Data Element": "calendarCode", "Reviewed?": "Yes"},
        ], generated=old)
        # A scores sidecar newer than the workbook stamp.
        (out / "az_scores_source.json").write_text("{}", encoding="utf-8")
        report = _ingest(wb, out, cur)
        assert report.stale_warning is not None
        assert "predates" in report.stale_warning
        cap = json.loads(
            state_curation_path("AZ", cur).read_text(encoding="utf-8")
        )["entries"]["AZ|Calendar|calendarCode"]["values"]["reviewed"]
        assert cap["workbook_generated"] == old


# ---------------------------------------------------------------------------
# Override disagreement (§8.4)
# ---------------------------------------------------------------------------


class TestOverrideDisagreements:
    def _scores(self, adjusted=2.5, base=2):
        return {
            "AZ|Calendar|calendarCode": {
                "entity": "Calendar",
                "element_name": "calendarCode",
                "adjusted_nachos_score": adjusted,
                "dimensions": {"nachos_score": {"value": base}},
            }
        }

    def test_disagreement_surfaces(self):
        cur = {"AZ|Calendar|calendarCode": {
            "analyst_adjusted_override": 1.0}}
        [d] = override_disagreements(cur, self._scores(), "source")
        assert d["axis"] == "adjusted"
        assert d["override"] == 1.0
        assert d["engine_value"] == 2.5
        assert d["why"] == "analyst override 1 vs engine adjusted 2.5"

    def test_base_disagreement_surfaces(self):
        """Issue #250 — the base axis compares against the rule-cascade
        tier at dimensions.nachos_score.value, never the adjusted."""
        cur = {"AZ|Calendar|calendarCode": {
            "analyst_base_override": 1.0}}
        [d] = override_disagreements(cur, self._scores(), "source")
        assert d["axis"] == "base"
        assert d["override"] == 1.0
        assert d["engine_value"] == 2.0
        assert d["why"] == "analyst base override 1 vs engine base 2"

    def test_both_axes_yield_one_row_per_axis(self):
        # A record contested on both layers renders two disagreements —
        # adjusted (the headline axis) first — so each why-text stays
        # unambiguous (the #248 clustering diagnostic splits on `axis`).
        cur = {"AZ|Calendar|calendarCode": {
            "analyst_adjusted_override": 1.5,
            "analyst_base_override": 1.0,
        }}
        adj, base = override_disagreements(cur, self._scores(), "source")
        assert (adj["axis"], base["axis"]) == ("adjusted", "base")
        assert adj["why"] == "analyst override 1.5 vs engine adjusted 2.5"
        assert base["why"] == "analyst base override 1 vs engine base 2"

    def test_agreement_is_silent(self):
        cur = {"AZ|Calendar|calendarCode": {
            "analyst_adjusted_override": 2.5,
            "analyst_base_override": 2.0,
        }}
        assert override_disagreements(cur, self._scores(), "source") == []

    def test_half_tier_base_override_compares(self):
        # Flag-don't-reject: a deliberate half-tier judgment (1.5 vs an
        # integer engine tier) participates in disagreement math.
        cur = {"AZ|Calendar|calendarCode": {
            "analyst_base_override": 1.5}}
        [d] = override_disagreements(cur, self._scores(), "source")
        assert d["why"] == "analyst base override 1.5 vs engine base 2"

    def test_no_engine_score_is_silent(self):
        cur = {"AZ|Other|thing": {"analyst_adjusted_override": 1.0}}
        assert override_disagreements(cur, self._scores(), "source") == []

    def test_missing_base_dimension_is_silent(self):
        # A sidecar without a nachos_score dimension block can't anchor
        # a base comparison; the adjusted axis is unaffected.
        scores = self._scores()
        scores["AZ|Calendar|calendarCode"]["dimensions"] = {}
        cur = {"AZ|Calendar|calendarCode": {
            "analyst_adjusted_override": 1.0,
            "analyst_base_override": 1.0,
        }}
        [d] = override_disagreements(cur, scores, "source")
        assert d["axis"] == "adjusted"

    def test_non_numeric_override_is_silent(self):
        assert override_disagreements(
            {"AZ|Calendar|calendarCode": {
                "analyst_adjusted_override": "not a number",
                "analyst_base_override": "tier 2, surely"}},
            self._scores(), "source",
        ) == []

    def test_numeric_string_override_still_compares(self):
        # A parseable string ("1.0") entered as text still participates.
        [d] = override_disagreements(
            {"AZ|Calendar|calendarCode": {
                "analyst_adjusted_override": "1.0"}},
            self._scores(), "source",
        )
        assert d["override"] == 1.0

    def test_no_override_column_is_silent(self):
        cur = {"AZ|Calendar|calendarCode": {"edfi_comments": "hello"}}
        assert override_disagreements(cur, self._scores(), "source") == []


# ---------------------------------------------------------------------------
# End-to-end round-trip: build → hand-edit → ingest → rebuild
# ---------------------------------------------------------------------------


class TestRoundTripEndToEnd:
    def test_edits_survive_regeneration(self, tmp_path, monkeypatch):
        from src.report import analyst
        from tests.test_report_analyst import _write_az_fixture

        _write_az_fixture(tmp_path)
        out = tmp_path / "data" / "out"
        cur = tmp_path / "data" / "curation"
        monkeypatch.setattr(analyst, "_OUT_DIR", out)
        monkeypatch.setattr(analyst, "_SPINE_DIR", tmp_path / "data" / "spine")
        monkeypatch.setattr(curation, "_OUT_DIR", out)
        monkeypatch.setattr(curation, "_CURATION_DIR", cur)

        # 1. Generate; the band is empty.
        [path] = analyst.run(state="AZ", out_dir=out)
        wb = openpyxl.load_workbook(path)
        ws = wb["Details"]
        headers = [c.value for c in ws[1]]
        rev_col = headers.index("Reviewed?") + 1
        com_col = headers.index("Ed-Fi Comments") + 1
        base_ovr_col = headers.index("Analyst Base Score (override)") + 1
        adj_col = headers.index("Adjusted NACHOS Score") + 1
        base_col = headers.index("Base NACHOS Score") + 1
        assert ws.cell(row=2, column=rev_col).value is None

        # 2. Analyst hand-edits three band cells (incl. a leading-= value
        # stored as TEXT, the way Excel stores it — the formula-corruption
        # guard must hold when the re-render writes it back — and a
        # base-axis override, issue #250).
        ws.cell(row=2, column=rev_col, value="Yes")
        com_cell = ws.cell(row=2, column=com_col,
                           value="=looks derived, confirm")
        com_cell.data_type = "s"
        ws.cell(row=2, column=base_ovr_col, value=1.5)
        engine_adj_before = ws.cell(row=2, column=adj_col).value
        engine_base_before = ws.cell(row=2, column=base_col).value
        row2_key_cells = [
            ws.cell(row=2, column=headers.index(h) + 1).value
            for h in ("State", "Entity Name", "Data Element")
        ]
        wb.save(path)

        # 3. Ingest.
        report = curation.run(path, author="maria")
        assert report.captured_per_column == {
            "reviewed": 1, "edfi_comments": 1, "analyst_base_override": 1,
        }

        # 4. Regenerate — the band repopulates; engine cells untouched
        # (the base override never mutates the engine's base tier).
        [path2] = analyst.run(state="AZ", out_dir=out)
        wb2 = openpyxl.load_workbook(path2)
        ws2 = wb2["Details"]
        headers2 = [c.value for c in ws2[1]]
        assert headers2 == headers
        assert ws2.cell(row=2, column=rev_col).value == "Yes"
        assert (
            ws2.cell(row=2, column=com_col).value
            == "=looks derived, confirm"
        )
        assert ws2.cell(row=2, column=base_ovr_col).value == 1.5
        assert ws2.cell(row=2, column=adj_col).value == engine_adj_before
        assert ws2.cell(row=2, column=base_col).value == engine_base_before
        assert [
            ws2.cell(row=2, column=headers2.index(h) + 1).value
            for h in ("State", "Entity Name", "Data Element")
        ] == row2_key_cells

        # 5. Ingesting the regenerated workbook is a no-op (values match).
        report2 = curation.run(path2)
        assert report2.captured_per_column == {}
        assert report2.unchanged == 3


class TestOverrideDisagreementsForState:
    """The loader-level wrapper gates on the CAPTURED lens (the flat
    view drops provenance, so the gate lives in the sidecar reader).
    Each override axis carries its own lens stamp and is gated
    independently (issue #250)."""

    def _write_payload(self, cur_dir, lens=None, values=None):
        if values is None:
            values = {"analyst_adjusted_override": {
                "value": 1.0, "lens": lens}}
        cur_dir.mkdir(parents=True, exist_ok=True)
        (cur_dir / "az.json").write_text(json.dumps({
            "version": 1,
            "state": "AZ",
            "entries": {
                "AZ|Calendar|calendarCode": {
                    "values": values,
                    "history": [],
                }
            },
        }), encoding="utf-8")

    _SCORES = {
        "AZ|Calendar|calendarCode": {
            "entity": "Calendar",
            "element_name": "calendarCode",
            "adjusted_nachos_score": 2.5,
            "dimensions": {"nachos_score": {"value": 2}},
        }
    }

    def test_same_lens_fires(self, tmp_path):
        from src.report.curation import override_disagreements_for

        self._write_payload(tmp_path, "source")
        [d] = override_disagreements_for(
            "AZ", self._SCORES, "source", base=tmp_path
        )
        assert d["override"] == 1.0

    def test_base_same_lens_fires(self, tmp_path):
        from src.report.curation import override_disagreements_for

        self._write_payload(tmp_path, values={
            "analyst_base_override": {"value": 1.0, "lens": "source"}})
        [d] = override_disagreements_for(
            "AZ", self._SCORES, "source", base=tmp_path
        )
        assert d["axis"] == "base"
        assert d["why"] == "analyst base override 1 vs engine base 2"

    def test_cross_lens_is_silent(self, tmp_path):
        from src.report.curation import override_disagreements_for

        self._write_payload(tmp_path, "spine")
        assert override_disagreements_for(
            "AZ", self._SCORES, "source", base=tmp_path
        ) == []

    def test_axes_gate_independently(self, tmp_path):
        # Adjusted captured on source, base captured on spine — only
        # the axis whose captured lens matches the render surfaces.
        from src.report.curation import override_disagreements_for

        self._write_payload(tmp_path, values={
            "analyst_adjusted_override": {"value": 1.0, "lens": "source"},
            "analyst_base_override": {"value": 1.0, "lens": "spine"},
        })
        [d] = override_disagreements_for(
            "AZ", self._SCORES, "source", base=tmp_path
        )
        assert d["axis"] == "adjusted"

    def test_missing_sidecar_is_silent(self, tmp_path):
        from src.report.curation import override_disagreements_for

        assert override_disagreements_for(
            "AZ", self._SCORES, "source", base=tmp_path
        ) == []


class TestTrackerIngest:
    """`review ingest` also reads the Commitment Tracker's analyst
    columns (same record key, same sidecar)."""

    def test_tracker_columns_captured(self, tmp_path):
        out = tmp_path / "out"
        cur = tmp_path / "curation"
        _write_elements(out, "AZ", [("Calendar", "calendarCode")])
        path = tmp_path / "az_analyst.xlsx"
        # Details sheet (required) + a Commitment Tracker sheet.
        _make_workbook(path, [
            {"State": "AZ", "Entity Name": "Calendar",
             "Data Element": "calendarCode", "Reviewed?": "Yes"},
        ])
        wb = openpyxl.load_workbook(path)
        ws = wb.create_sheet("Commitment Tracker")
        ws.append(["Row #", "State", "Entity Name", "Data Element",
                   "NACHOS Points Resolved", "Adoption Timeline",
                   "Commitment Status", "Comments"])
        ws.append([5, "AZ", "Calendar", "calendarCode", 3,
                   "2027-28 school year", "Committed", "state confirmed"])
        # Subtotal/Grand Total rows carry no identity — must be ignored.
        ws.append([None, None, "Calendar — subtotal", None, 3,
                   None, None, None])
        wb.save(path)

        report = _ingest(path, out, cur)
        values = curation_values_for("AZ", cur)["AZ|Calendar|calendarCode"]
        assert values["reviewed"] == "Yes"
        assert values["adoption_timeline"] == "2027-28 school year"
        assert values["commitment_status"] == "Committed"
        assert values["commitment_comments"] == "state confirmed"
        assert report.captured_per_column == {
            "reviewed": 1, "adoption_timeline": 1,
            "commitment_status": 1, "commitment_comments": 1,
        }


# ---------------------------------------------------------------------------
# Adjudication overlay (issue #248 Part B — curation schema v2)
# ---------------------------------------------------------------------------


def _write_scores(out: Path, state: str = "AZ", lens: str = "source",
                  adjusted: float | None = 2.5) -> None:
    out.mkdir(parents=True, exist_ok=True)
    (out / f"{state.lower()}_scores_{lens}.json").write_text(
        json.dumps({
            "state": state,
            "lens": lens,
            "scores": [{
                "record_key": f"{state}|Calendar|calendarCode",
                "entity": "Calendar",
                "element_name": "calendarCode",
                "adjusted_nachos_score": adjusted,
                "dimensions": {"nachos_score": {"value": 2}},
            }],
        }),
        encoding="utf-8",
    )


def _adjudicate(out, cur, **kw):
    kw.setdefault("value", 2.0)
    kw.setdefault("agreed_by", ("Doug", "Maria", "Chris"))
    kw.setdefault("rationale", "team consensus")
    return curation.adjudicate(
        "az", "Calendar", "calendarCode",
        out_base=out, curation_base=cur, **kw,
    )


class TestAdjudicate:
    """The `mc review adjudicate` writer — schema v2 block with full
    provenance; the engine score is never touched."""

    def _env(self, tmp_path):
        out = tmp_path / "out"
        cur = tmp_path / "curation"
        _write_elements(out, "AZ", [("Calendar", "calendarCode"),
                                    ("School", "schoolId")])
        _write_scores(out)
        return out, cur

    def test_block_written_with_full_provenance(self, tmp_path, monkeypatch):
        from src.score.aggregate import SCORING_PLAN_VERSION

        out, cur = self._env(tmp_path)
        monkeypatch.setattr(
            curation, "_now_iso", lambda: "2026-07-13T12:00:00+00:00"
        )
        report = _adjudicate(out, cur)
        payload = json.loads(
            state_curation_path("AZ", cur).read_text(encoding="utf-8")
        )
        # Deliberate-bump tripwire: v3 = issue #249 fact-correction
        # blocks (v2 was the #248 adjudication block).
        assert payload["version"] == 3
        block = payload["entries"]["AZ|Calendar|calendarCode"]["adjudication"]
        assert block == {
            "status": "adjudicated",
            "axis": "adjusted",
            "value": 2.0,
            "lens": "source",
            "agreed_by": ["Doug", "Maria", "Chris"],
            "decided_at": "2026-07-13T12:00:00+00:00",
            "rationale": "team consensus",
            "engine_score_at_decision": 2.5,
            "plan_version_at_decision": SCORING_PLAN_VERSION,
        }
        assert report.engine_score_at_decision == 2.5
        assert report.replaced_prior is False
        assert report.written_path is not None
        # Engine sidecar byte-untouched.
        scores = json.loads(
            (out / "az_scores_source.json").read_text(encoding="utf-8")
        )
        assert scores["scores"][0]["adjusted_nachos_score"] == 2.5

    def test_unknown_record_key_aborts_nothing_written(self, tmp_path):
        out, cur = self._env(tmp_path)
        with pytest.raises(curation.AdjudicationAbort, match="unknown record key"):
            curation.adjudicate(
                "az", "Calendar", "notAColumn",
                value=2.0, agreed_by=("Doug",), rationale="x",
                out_base=out, curation_base=cur,
            )
        assert not state_curation_path("AZ", cur).exists()

    def test_missing_scores_sidecar_aborts(self, tmp_path):
        out = tmp_path / "out"
        cur = tmp_path / "curation"
        _write_elements(out, "AZ", [("Calendar", "calendarCode")])
        with pytest.raises(curation.AdjudicationAbort, match="unscored record"):
            _adjudicate(out, cur)

    def test_unscored_record_aborts(self, tmp_path):
        out, cur = self._env(tmp_path)
        _write_scores(out, adjusted=None)
        with pytest.raises(curation.AdjudicationAbort, match="unscored record"):
            _adjudicate(out, cur)

    def test_no_agreed_by_aborts(self, tmp_path):
        out, cur = self._env(tmp_path)
        with pytest.raises(curation.AdjudicationAbort, match="agreed-by"):
            _adjudicate(out, cur, agreed_by=("", "  "))

    def test_readjudication_pushes_prior_block_to_history(self, tmp_path):
        out, cur = self._env(tmp_path)
        _adjudicate(out, cur, value=2.0)
        report = _adjudicate(out, cur, value=1.5, rationale="revised")
        assert report.replaced_prior is True
        entry = json.loads(
            state_curation_path("AZ", cur).read_text(encoding="utf-8")
        )["entries"]["AZ|Calendar|calendarCode"]
        assert entry["adjudication"]["value"] == 1.5
        [hist] = entry["history"]
        assert hist["column"] == "adjudication"
        assert hist["value"]["value"] == 2.0

    def test_dry_run_writes_nothing(self, tmp_path):
        out, cur = self._env(tmp_path)
        report = _adjudicate(out, cur, dry_run=True)
        assert report.dry_run is True
        assert report.written_path is None
        assert not state_curation_path("AZ", cur).exists()

    def test_coexists_with_values_and_survives_ingest(self, tmp_path):
        """The round-trip promise: ingest only merges `values`/`history`,
        so an adjudication block survives a later workbook ingest, and
        the flat re-apply view never exposes it."""
        out, cur = self._env(tmp_path)
        _adjudicate(out, cur)
        wb_path = tmp_path / "az_analyst.xlsx"
        _make_workbook(wb_path, [
            {"State": "AZ", "Entity Name": "Calendar",
             "Data Element": "calendarCode", "Reviewed?": "Yes"},
        ])
        ingest_workbook(wb_path, out_base=out, curation_base=cur)
        entry = json.loads(
            state_curation_path("AZ", cur).read_text(encoding="utf-8")
        )["entries"]["AZ|Calendar|calendarCode"]
        assert entry["adjudication"]["value"] == 2.0
        assert entry["values"]["reviewed"]["value"] == "Yes"
        assert curation_values_for("AZ", cur)[
            "AZ|Calendar|calendarCode"
        ] == {"reviewed": "Yes"}


class TestAdjudicationsFor:
    """The render-side reader: lens gate + freshness (engine epsilon,
    plan version, record-gone), computed at read time."""

    _KEY = "AZ|Calendar|calendarCode"

    def _write_block(self, cur_dir, **overrides):
        from src.score.aggregate import SCORING_PLAN_VERSION

        block = {
            "status": "adjudicated",
            "axis": "adjusted",
            "value": 2.0,
            "lens": "source",
            "agreed_by": ["Doug", "Maria"],
            "decided_at": "2026-07-13T12:00:00+00:00",
            "rationale": "team consensus",
            "engine_score_at_decision": 2.5,
            "plan_version_at_decision": SCORING_PLAN_VERSION,
        }
        block.update(overrides)
        cur_dir.mkdir(parents=True, exist_ok=True)
        (cur_dir / "az.json").write_text(json.dumps({
            "version": 2,
            "state": "AZ",
            "entries": {self._KEY: {"values": {}, "history": [],
                                    "adjudication": block}},
        }), encoding="utf-8")

    _SCORES = {
        "AZ|Calendar|calendarCode": {
            "entity": "Calendar",
            "element_name": "calendarCode",
            "adjusted_nachos_score": 2.5,
            "dimensions": {"nachos_score": {"value": 2}},
        }
    }

    def test_fresh_when_engine_and_plan_match(self, tmp_path):
        self._write_block(tmp_path)
        adj = curation.adjudications_for(
            "AZ", self._SCORES, "source", base=tmp_path
        )[self._KEY]
        assert adj["fresh"] is True
        assert adj["stale_reason"] is None
        assert adj["value"] == 2.0
        assert adj["engine_value"] == 2.5
        assert adj["entity"] == "Calendar"

    def test_fresh_within_epsilon(self, tmp_path):
        self._write_block(tmp_path, engine_score_at_decision=2.501)
        adj = curation.adjudications_for(
            "AZ", self._SCORES, "source", base=tmp_path
        )[self._KEY]
        assert adj["fresh"] is True

    def test_agreeing_adjudication_is_still_fresh(self, tmp_path):
        # Consensus that AGREES with the engine still renders — blank
        # means "none or stale" only, and the stamp re-arms staleness.
        self._write_block(tmp_path, value=2.5)
        adj = curation.adjudications_for(
            "AZ", self._SCORES, "source", base=tmp_path
        )[self._KEY]
        assert adj["fresh"] is True

    def test_stale_when_engine_moved(self, tmp_path):
        self._write_block(tmp_path, engine_score_at_decision=0.5)
        adj = curation.adjudications_for(
            "AZ", self._SCORES, "source", base=tmp_path
        )[self._KEY]
        assert adj["fresh"] is False
        assert adj["stale_reason"] == "engine adjusted moved 0.5 → 2.5"

    def test_stale_when_plan_bumped(self, tmp_path):
        from src.score.aggregate import SCORING_PLAN_VERSION

        self._write_block(tmp_path, plan_version_at_decision="0")
        adj = curation.adjudications_for(
            "AZ", self._SCORES, "source", base=tmp_path
        )[self._KEY]
        assert adj["fresh"] is False
        assert adj["stale_reason"] == f"plan v0 → v{SCORING_PLAN_VERSION}"

    def test_stale_when_record_gone(self, tmp_path):
        self._write_block(tmp_path)
        adj = curation.adjudications_for(
            "AZ", {}, "source", base=tmp_path
        )[self._KEY]
        assert adj["fresh"] is False
        assert adj["stale_reason"] == (
            "record no longer scored at the current plan"
        )
        # Identity still resolves from the record key for the register.
        assert adj["entity"] == "Calendar"
        assert adj["element_name"] == "calendarCode"

    def test_cross_lens_is_invisible(self, tmp_path):
        self._write_block(tmp_path, lens="spine")
        assert curation.adjudications_for(
            "AZ", self._SCORES, "source", base=tmp_path
        ) == {}

    def test_non_adjusted_axis_skipped(self, tmp_path):
        # Forward-compat: a future base-axis block must not surface
        # through today's adjusted-only reader.
        self._write_block(tmp_path, axis="base")
        assert curation.adjudications_for(
            "AZ", self._SCORES, "source", base=tmp_path
        ) == {}

    def test_missing_sidecar_is_silent(self, tmp_path):
        assert curation.adjudications_for(
            "AZ", self._SCORES, "source", base=tmp_path
        ) == {}


# ---------------------------------------------------------------------------
# Fact-level curation (issue #249 — schema v3)
# ---------------------------------------------------------------------------


def _write_fact_scores(
    out: Path,
    state: str = "AZ",
    lens: str = "source",
    fact_provenance: dict | None = None,
) -> None:
    """Scores sidecar with a fact_provenance block on the one record."""
    if fact_provenance is None:
        fact_provenance = {
            "has_conditional_logic": {
                "value": True,
                "confidence": "medium",
                "downgraded": False,
                "downgrade_reason": None,
            },
            "semantic_class": {
                "value": "aligned",
                "confidence": "high",
                "downgraded": False,
                "downgrade_reason": None,
            },
        }
    out.mkdir(parents=True, exist_ok=True)
    (out / f"{state.lower()}_scores_{lens}.json").write_text(
        json.dumps({
            "state": state,
            "lens": lens,
            "scores": [{
                "record_key": f"{state}|Calendar|calendarCode",
                "entity": "Calendar",
                "element_name": "calendarCode",
                "adjusted_nachos_score": 2.5,
                "dimensions": {"nachos_score": {"value": 2}},
                "fact_provenance": fact_provenance,
            }],
        }),
        encoding="utf-8",
    )


def _correct(out, cur, **kw):
    kw.setdefault("fact", "has_conditional_logic")
    kw.setdefault("value", "false")
    kw.setdefault("rationale", "calendar language, not conditionality")
    kw.setdefault("author", "Chris Moffatt")
    return curation.correct_fact(
        "az", "Calendar", "calendarCode",
        out_base=out, curation_base=cur, **kw,
    )


class TestCorrectFact:
    """The `mc review correct-fact` writer — schema v3 `facts` block
    with full provenance; the engine sidecar is never touched (the
    overlay applies at aggregate time)."""

    def _env(self, tmp_path, **scores_kw):
        out = tmp_path / "out"
        cur = tmp_path / "curation"
        _write_elements(out, "AZ", [("Calendar", "calendarCode"),
                                    ("School", "schoolId")])
        _write_fact_scores(out, **scores_kw)
        return out, cur

    def test_block_written_with_full_provenance(self, tmp_path, monkeypatch):
        from src.score.aggregate import SCORING_PLAN_VERSION

        out, cur = self._env(tmp_path)
        monkeypatch.setattr(
            curation, "_now_iso", lambda: "2026-07-13T12:00:00+00:00"
        )
        report = _correct(out, cur)
        payload = json.loads(
            state_curation_path("AZ", cur).read_text(encoding="utf-8")
        )
        assert payload["version"] == 3
        block = payload["entries"]["AZ|Calendar|calendarCode"]["facts"][
            "has_conditional_logic"
        ]
        assert block == {
            "value": False,
            "lens": "source",
            "author": "Chris Moffatt",
            "corrected_at": "2026-07-13T12:00:00+00:00",
            "rationale": "calendar language, not conditionality",
            "prior_value": True,
            "prior_provenance": "llm",
            "plan_version_at_correction": SCORING_PLAN_VERSION,
        }
        assert report.value is False
        assert report.prior_value is True
        assert report.replaced_prior is False
        assert report.written_path is not None
        assert "aggregate" in report.render_text()
        # Engine sidecar byte-untouched — the overlay is aggregate-time.
        scores = json.loads(
            (out / "az_scores_source.json").read_text(encoding="utf-8")
        )
        prov = scores["scores"][0]["fact_provenance"]
        assert prov["has_conditional_logic"]["value"] is True

    def test_unknown_record_key_aborts_nothing_written(self, tmp_path):
        out, cur = self._env(tmp_path)
        with pytest.raises(
            curation.FactCorrectionAbort, match="unknown record key"
        ):
            curation.correct_fact(
                "az", "Calendar", "notAColumn",
                fact="has_conditional_logic", value="false",
                rationale="x", author="Chris",
                out_base=out, curation_base=cur,
            )
        assert not state_curation_path("AZ", cur).exists()

    def test_unscored_record_aborts(self, tmp_path):
        out = tmp_path / "out"
        cur = tmp_path / "curation"
        _write_elements(out, "AZ", [("Calendar", "calendarCode")])
        with pytest.raises(
            curation.FactCorrectionAbort, match="unscored record"
        ):
            _correct(out, cur)

    def test_deterministic_fact_rejected(self, tmp_path):
        out, cur = self._env(tmp_path)
        with pytest.raises(
            curation.FactCorrectionAbort,
            match="deterministic fact.*code bug",
        ):
            _correct(out, cur, fact="definition_present", value="true")
        assert not state_curation_path("AZ", cur).exists()

    def test_unknown_fact_rejected(self, tmp_path):
        out, cur = self._env(tmp_path)
        with pytest.raises(
            curation.FactCorrectionAbort, match="unknown fact"
        ):
            _correct(out, cur, fact="not_a_fact", value="true")

    def test_fact_not_extracted_aborts(self, tmp_path):
        out, cur = self._env(
            tmp_path,
            fact_provenance={
                "semantic_class": {
                    "value": "aligned", "confidence": "high",
                    "downgraded": False, "downgrade_reason": None,
                },
            },
        )
        with pytest.raises(
            curation.FactCorrectionAbort, match="was not extracted"
        ):
            _correct(out, cur)

    def test_source_filtered_fact_aborts(self, tmp_path):
        out, cur = self._env(
            tmp_path,
            fact_provenance={
                "extension_is_necessary": {
                    "value": None, "confidence": "low",
                    "downgraded": True,
                    "downgrade_reason": "filtered_by_source",
                },
            },
        )
        with pytest.raises(
            curation.FactCorrectionAbort, match="source-filtered"
        ):
            _correct(out, cur, fact="extension_is_necessary", value="true")

    def test_missing_from_artifact_aborts_with_spine_borrow_hint(
        self, tmp_path
    ):
        out = tmp_path / "out"
        cur = tmp_path / "curation"
        _write_elements(out, "AZ", [("Calendar", "calendarCode")],
                        lens="spine")
        _write_fact_scores(
            out, lens="spine",
            fact_provenance={
                "extension_is_necessary": {
                    "value": None, "confidence": "low",
                    "downgraded": True,
                    "downgrade_reason": "missing_from_artifact",
                },
            },
        )
        with pytest.raises(
            curation.FactCorrectionAbort, match="SOURCE"
        ):
            curation.correct_fact(
                "az", "Calendar", "calendarCode",
                fact="extension_is_necessary", value="true",
                rationale="x", author="Chris", lens="spine",
                out_base=out, curation_base=cur,
            )

    def test_bool_value_validated(self, tmp_path):
        out, cur = self._env(tmp_path)
        with pytest.raises(
            curation.FactCorrectionAbort, match="bool fact"
        ):
            _correct(out, cur, value="maybe")

    def test_enum_value_validated(self, tmp_path):
        out, cur = self._env(tmp_path)
        with pytest.raises(
            curation.FactCorrectionAbort, match="not one of"
        ):
            _correct(out, cur, fact="semantic_class", value="sideways")

    def test_enum_value_accepted(self, tmp_path):
        out, cur = self._env(tmp_path)
        report = _correct(out, cur, fact="semantic_class",
                          value="divergent")
        assert report.value == "divergent"
        assert report.prior_value == "aligned"

    def test_noop_equal_value_aborts(self, tmp_path):
        out, cur = self._env(tmp_path)
        with pytest.raises(
            curation.FactCorrectionAbort, match="nothing to correct"
        ):
            _correct(out, cur, value="true")

    def test_downgraded_fact_is_correctable(self, tmp_path):
        """A downgraded (None-valued) fact is a prime correction target
        — asserting the value the cascade currently defaults to is NOT
        a no-op, because it clears the downgrade at overlay time."""
        out, cur = self._env(
            tmp_path,
            fact_provenance={
                "has_conditional_logic": {
                    "value": None, "confidence": "low",
                    "downgraded": True,
                    "downgrade_reason": "hallucinated_span",
                },
            },
        )
        report = _correct(out, cur, value="false")
        assert report.prior_value is None
        assert report.prior_provenance == "llm_downgraded:hallucinated_span"

    def test_recorrection_moves_prior_to_history(self, tmp_path):
        out, cur = self._env(tmp_path)
        _correct(out, cur, value="false")
        report = _correct(out, cur, value="true",
                          rationale="second look: the gate is real",
                          author="Maria")
        assert report.replaced_prior is True
        payload = json.loads(
            state_curation_path("AZ", cur).read_text(encoding="utf-8")
        )
        entry = payload["entries"]["AZ|Calendar|calendarCode"]
        assert entry["facts"]["has_conditional_logic"]["value"] is True
        assert entry["facts"]["has_conditional_logic"]["author"] == "Maria"
        hist = [h for h in entry["history"]
                if h["column"] == "fact:has_conditional_logic"]
        assert len(hist) == 1
        assert hist[0]["value"]["value"] is False

    def test_blank_author_or_rationale_aborts(self, tmp_path):
        out, cur = self._env(tmp_path)
        with pytest.raises(
            curation.FactCorrectionAbort, match="--author"
        ):
            _correct(out, cur, author="   ")
        with pytest.raises(
            curation.FactCorrectionAbort, match="--rationale"
        ):
            _correct(out, cur, rationale="")

    def test_dry_run_writes_nothing(self, tmp_path):
        out, cur = self._env(tmp_path)
        report = _correct(out, cur, dry_run=True)
        assert report.dry_run is True
        assert report.written_path is None
        assert not state_curation_path("AZ", cur).exists()

    def test_coexists_with_values_and_adjudication(self, tmp_path):
        out, cur = self._env(tmp_path)
        _correct(out, cur)
        _adjudicate(out, cur)
        payload = json.loads(
            state_curation_path("AZ", cur).read_text(encoding="utf-8")
        )
        entry = payload["entries"]["AZ|Calendar|calendarCode"]
        assert "facts" in entry
        assert "adjudication" in entry
        assert entry["facts"]["has_conditional_logic"]["value"] is False

    def test_multiple_facts_one_entry(self, tmp_path):
        out, cur = self._env(tmp_path)
        _correct(out, cur)
        _correct(out, cur, fact="semantic_class", value="divergent")
        payload = json.loads(
            state_curation_path("AZ", cur).read_text(encoding="utf-8")
        )
        facts = payload["entries"]["AZ|Calendar|calendarCode"]["facts"]
        assert set(facts) == {"has_conditional_logic", "semantic_class"}


class TestFactCorrectionsFor:
    """The aggregate-overlay reader — same-lens blocks, verbatim."""

    _KEY = "AZ|Calendar|calendarCode"

    def _write_block(self, base: Path, lens: str = "source") -> None:
        payload = {
            "version": 3,
            "state": "AZ",
            "updated_at": "2026-07-13T12:00:00+00:00",
            "entries": {
                self._KEY: {
                    "values": {},
                    "history": [],
                    "facts": {
                        "has_conditional_logic": {
                            "value": False,
                            "lens": lens,
                            "author": "Chris Moffatt",
                            "corrected_at": "2026-07-13T12:00:00+00:00",
                            "rationale": "calendar, not conditionality",
                            "prior_value": True,
                            "prior_provenance": "llm",
                            "plan_version_at_correction": "28",
                        },
                    },
                },
            },
        }
        base.mkdir(parents=True, exist_ok=True)
        (base / "az.json").write_text(
            json.dumps(payload), encoding="utf-8"
        )

    def test_same_lens_blocks_returned_verbatim(self, tmp_path):
        self._write_block(tmp_path)
        out = curation.fact_corrections_for("AZ", "source", base=tmp_path)
        assert set(out) == {self._KEY}
        block = out[self._KEY]["has_conditional_logic"]
        assert block["value"] is False
        assert block["rationale"] == "calendar, not conditionality"

    def test_cross_lens_is_invisible(self, tmp_path):
        self._write_block(tmp_path, lens="spine")
        assert curation.fact_corrections_for(
            "AZ", "source", base=tmp_path
        ) == {}

    def test_entries_without_facts_are_skipped(self, tmp_path):
        payload = {
            "version": 3, "state": "AZ", "updated_at": None,
            "entries": {self._KEY: {"values": {}, "history": []}},
        }
        tmp_path.mkdir(parents=True, exist_ok=True)
        (tmp_path / "az.json").write_text(
            json.dumps(payload), encoding="utf-8"
        )
        assert curation.fact_corrections_for(
            "AZ", "source", base=tmp_path
        ) == {}

    def test_missing_sidecar_is_silent(self, tmp_path):
        assert curation.fact_corrections_for(
            "AZ", "source", base=tmp_path
        ) == {}

    def test_v2_files_stay_readable(self, tmp_path):
        """Additive-convergent contract: a pre-#249 (v2) sidecar simply
        has no `facts` blocks."""
        payload = {
            "version": 2, "state": "AZ", "updated_at": None,
            "entries": {self._KEY: {"values": {"reviewed": {
                "value": "Yes", "lens": "source",
            }}, "history": []}},
        }
        tmp_path.mkdir(parents=True, exist_ok=True)
        (tmp_path / "az.json").write_text(
            json.dumps(payload), encoding="utf-8"
        )
        assert curation.fact_corrections_for(
            "AZ", "source", base=tmp_path
        ) == {}
        assert curation.curation_values_for("AZ", base=tmp_path)

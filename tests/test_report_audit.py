"""Tests for the on-demand audit workbook (`mc report audit`, seq 4 PR C)."""

from __future__ import annotations

import inspect

import click
import openpyxl

from src.report import analyst, audit
from tests.test_report_analyst import (
    TestPhaseDScoringIntegration,
    _az_elements,
    _write_az_fixture,
)


class TestCliWiring:
    def test_run_is_plain_function_not_click_command(self):
        assert not isinstance(audit.run, click.Command)
        assert inspect.isfunction(audit.run)

    def test_cli_command_calls_module_run(self, monkeypatch, tmp_path):
        from click.testing import CliRunner

        from src.cli import cli

        calls: dict = {}

        def fake_run(*, state, lens="source", out=None):
            calls.update(state=state, lens=lens)
            p = tmp_path / "az_audit.xlsx"
            p.write_text("x", encoding="utf-8")
            return p

        monkeypatch.setattr(audit, "run", fake_run)
        result = CliRunner().invoke(
            cli, ["report", "audit", "--state", "az", "--lens", "spine"]
        )
        assert result.exit_code == 0, result.output
        assert calls == {"state": "AZ", "lens": "spine"}


class TestAuditWorkbook:
    def _setup(self, tmp_path, monkeypatch, *, with_scores=True):
        _write_az_fixture(tmp_path)
        out = tmp_path / "data" / "out"
        if with_scores:
            TestPhaseDScoringIntegration()._write_scores_sidecar(
                out, "AZ", "source", _az_elements().elements
            )
        monkeypatch.setattr(analyst, "_OUT_DIR", out)
        monkeypatch.setattr(analyst, "_SPINE_DIR", tmp_path / "data" / "spine")
        return out

    def test_writes_audit_workbook_with_expected_shape(
        self, tmp_path, monkeypatch
    ):
        out = self._setup(tmp_path, monkeypatch)
        path = audit.run(state="AZ", lens="source", out=out)
        assert path.name == "az_audit.xlsx"
        wb = openpyxl.load_workbook(path)
        assert wb.sheetnames == ["Readme", "Audit Trail", "Legend"]
        ws = wb["Audit Trail"]
        # Row # + the 63-column source factory output (belt-and-braces
        # width pin mirroring the old in-workbook assertion).
        assert ws.max_column == 64
        assert ws.max_row == 3  # 2 fixture records + header
        # Registry finish applies (audit tab color + freeze).
        assert ws.freeze_panes == "G2"
        assert ws.sheet_properties.tabColor.rgb == "FFD9D9D9"

    def test_spine_lens_suffix_and_width(self, tmp_path, monkeypatch):
        _write_az_fixture(tmp_path)
        out = tmp_path / "data" / "out"
        # Spine artifact = reuse the source fixture rows under the spine
        # filename (the audit builder only needs a valid elements file).
        (out / "az_elements_spine.json").write_text(
            (out / "az_elements_source.json").read_text(encoding="utf-8"),
            encoding="utf-8",
        )
        TestPhaseDScoringIntegration()._write_scores_sidecar(
            out, "AZ", "spine", _az_elements().elements
        )
        monkeypatch.setattr(analyst, "_OUT_DIR", out)
        monkeypatch.setattr(analyst, "_SPINE_DIR", tmp_path / "data" / "spine")
        path = audit.run(state="AZ", lens="spine", out=out)
        assert path.name == "az_audit_spine.xlsx"
        ws = openpyxl.load_workbook(path)["Audit Trail"]
        assert ws.max_column == 67  # Row # + 66 spine factory columns

    def test_unscored_state_still_renders_full_rows(
        self, tmp_path, monkeypatch
    ):
        out = self._setup(tmp_path, monkeypatch, with_scores=False)
        path = audit.run(state="AZ", lens="source", out=out)
        ws = openpyxl.load_workbook(path)["Audit Trail"]
        # The audit surface must not shrink: every row present even
        # without a sidecar (fact/dim cells blank).
        assert ws.max_row == 3

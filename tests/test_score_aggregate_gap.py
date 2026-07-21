"""Spine-anchored gap aggregate — issue #73 Step 1.

Covers the deterministic-only sidecar pass over ``{state}_elements_gap.json``:

- ``discovery_lens="spine_anchored"`` lands on every emitted record.
- Deterministic structural facts populate (``fk_chain_depth``,
  ``reference_fan_out``, ``sub_collection_depth``,
  ``entity_extension_footprint``, ``descriptor_enum_breadth``).
- LLM-dependent facts surface as downgraded (no artifact on disk for
  Step 1) — the rule cascade lands them at tier 0 with low confidence.
- Sidecar header carries the ``lens="spine_gap"`` discriminator and
  echoes the surfacer envelope's ``generated_at`` / discovery counts.
- CLI module-level ``run()`` is a plain function (NOT @click.command).
"""

from __future__ import annotations

import json
from datetime import datetime, timezone
from pathlib import Path

from src.models.edfi_catalog import (
    EdFiCatalog,
    EntityEntry,
    ExtensionEntry,
    PropertyInfo,
    ReferenceInfo,
)
from src.models.spine import SpineSourceURLs, StateSpine
from src.score.aggregate_gap import (
    DISCOVERY_LENS_SPINE_ANCHORED,
    GAP_LENS_TAG,
    run as run_gap,
)


def _spine(entities: dict, extensions: dict | None = None) -> StateSpine:
    catalog = EdFiCatalog(
        version="4.0",
        entity_count=len(entities),
        extension_count=len(extensions or {}),
        entities=entities,
        extensions=extensions or {},
    )
    return StateSpine(
        state="TX",
        edfi_version="4.0",
        fetched_at=datetime.now(tz=timezone.utc),
        source_urls=SpineSourceURLs(resources="http://test/resources.json"),
        catalog=catalog,
    )


def _write_gap_artifact(path: Path, gaps: list[dict]) -> None:
    payload = {
        "state": "TX",
        "generated_at": "2026-04-29T00:00:00Z",
        "gap_count": len(gaps),
        "discovery_counts": {},
        "gaps": gaps,
    }
    for g in gaps:
        d = g.get("discovery", "unknown")
        payload["discovery_counts"][d] = payload["discovery_counts"].get(d, 0) + 1
    path.write_text(json.dumps(payload, indent=2), encoding="utf-8")


def _read(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


# ---------------------------------------------------------------------------
# Sidecar shape
# ---------------------------------------------------------------------------


def test_emits_discovery_lens_spine_anchored_on_every_record(tmp_path: Path) -> None:
    """Step 1 sanity: every gap-derived record carries the discovery_lens flag."""
    spine = _spine(
        entities={"Calendar": EntityEntry(properties={"calendarCode": PropertyInfo()})}
    )
    spine_path = tmp_path / "tx_spine.json"
    spine_path.write_text(spine.model_dump_json())
    gap_path = tmp_path / "tx_elements_gap.json"
    _write_gap_artifact(gap_path, [
        {
            "state": "TX",
            "entity": "Calendar",
            "element_name": "calendarCode",
            "discovery": "spine_within_documented_entity",
            "documented_in_source": False,
            "spine_data_type": "String",
            "spine_extension_name": None,
            "rationale": "x",
        },
    ])
    out_path = tmp_path / "tx_scores_gap.json"

    header = run_gap(
        state="TX",
        gap_path=gap_path,
        spine_path=spine_path,
        out_path=out_path,
    )
    payload = _read(out_path)

    assert header["lens"] == GAP_LENS_TAG
    assert payload["lens"] == GAP_LENS_TAG
    assert payload["record_count"] == 1
    record = payload["scores"][0]
    assert record["discovery_lens"] == DISCOVERY_LENS_SPINE_ANCHORED


def test_deterministic_structural_facts_populate(tmp_path: Path) -> None:
    """fk_chain_depth, reference_fan_out, sub_collection_depth, descriptor_enum_breadth,
    entity_extension_footprint should all land on the FactView with high confidence
    even though no LLM extracts ran."""
    # Two-entity catalog: Student → School (FK chain depth 1 on Student),
    # School with one extension → entity_extension_footprint 1.
    spine = _spine(
        entities={
            "School": EntityEntry(
                properties={"schoolId": PropertyInfo(is_identity=True)},
            ),
            "Student": EntityEntry(
                properties={"studentUniqueId": PropertyInfo(is_identity=True)},
                references={
                    "schoolReference": ReferenceInfo(
                        entity="School",
                        key_properties={"schoolId": PropertyInfo(is_identity=True)},
                    )
                },
            ),
        },
        extensions={
            "tx_schoolExtension": ExtensionEntry(
                extends_entity="School",
                source_prefix="tx",
                properties={"tx_extra": PropertyInfo()},
            )
        },
    )
    spine_path = tmp_path / "tx_spine.json"
    spine_path.write_text(spine.model_dump_json())
    gap_path = tmp_path / "tx_elements_gap.json"
    _write_gap_artifact(gap_path, [
        # A gap row on Student — should pick up fk_chain_depth=1 (Student → School).
        {
            "state": "TX", "entity": "Student", "element_name": "firstName",
            "discovery": "spine_within_documented_entity",
            "documented_in_source": False,
            "spine_data_type": "String",
            "spine_extension_name": None,
            "rationale": "x",
        },
        # A gap row on School — should pick up entity_extension_footprint=1
        # and reference_fan_out=1 (Student references School).
        {
            "state": "TX", "entity": "School", "element_name": "shortNameOfInstitution",
            "discovery": "spine_within_documented_entity",
            "documented_in_source": False,
            "spine_data_type": "String",
            "spine_extension_name": None,
            "rationale": "x",
        },
    ])
    out_path = tmp_path / "tx_scores_gap.json"

    run_gap(state="TX", gap_path=gap_path, spine_path=spine_path, out_path=out_path)
    payload = _read(out_path)

    by_entity = {s["entity"]: s for s in payload["scores"]}
    student_inputs = by_entity["Student"]["dimensions"]["structural_depth"]["inputs_used"]
    school_inputs = by_entity["School"]["dimensions"]["structural_depth"]["inputs_used"]

    assert student_inputs["fk_chain_depth"] == 1
    assert school_inputs["reference_fan_out"] == 1
    assert school_inputs["entity_extension_footprint"] == 1
    assert student_inputs["sub_collection_depth"] == 0
    # Deterministic facts should come back high-confidence.
    sd_conf = by_entity["Student"]["dimensions"]["structural_depth"]["confidence"]
    assert sd_conf == "high"


def test_llm_facts_surface_as_downgraded(tmp_path: Path) -> None:
    """Step 1: LLM rule inputs land as downgraded with confidence=low.

    Asserts the contract that step 2 / step 3 will lift these into real
    values — until then the deterministic-only pass surfaces them so
    consumers see exactly which facts are pending extraction.
    """
    spine = _spine(
        entities={"Calendar": EntityEntry(properties={"calendarCode": PropertyInfo()})}
    )
    spine_path = tmp_path / "tx_spine.json"
    spine_path.write_text(spine.model_dump_json())
    gap_path = tmp_path / "tx_elements_gap.json"
    _write_gap_artifact(gap_path, [
        {
            "state": "TX", "entity": "Calendar", "element_name": "calendarCode",
            "discovery": "spine_within_documented_entity",
            "documented_in_source": False,
            "spine_data_type": "String",
            "spine_extension_name": None,
            "rationale": "x",
        },
    ])
    out_path = tmp_path / "tx_scores_gap.json"

    run_gap(state="TX", gap_path=gap_path, spine_path=spine_path, out_path=out_path)
    record = _read(out_path)["scores"][0]
    fp = record["fact_provenance"]
    # Spot-check three LLM rule inputs.
    for fact in ("definition_is_implementable", "has_conditional_logic", "documentation_style"):
        entry = fp.get(fact)
        assert entry is not None
        assert entry["downgraded"] is True
        assert entry["confidence"] == "low"
        assert entry["value"] is None


def test_extension_gap_row_marked_extension_source(tmp_path: Path) -> None:
    """Gap rows whose spine slot lives on a state extension carry source='extension'.

    Ensures the rule cascade's extension-aware branches behave the same
    way they would for a normal source-lens extension row — important
    so the conservative ``extension_necessity_unresolved`` adjustment
    fires on Step 1 (and disappears in Step 3 once the LLM extract
    populates ``extension_is_necessary``).
    """
    spine = _spine(
        entities={"Calendar": EntityEntry(properties={"calendarCode": PropertyInfo()})},
        extensions={
            "tx_calendarExtension": ExtensionEntry(
                extends_entity="Calendar",
                source_prefix="tx",
                properties={"tx_localCode": PropertyInfo()},
            )
        },
    )
    spine_path = tmp_path / "tx_spine.json"
    spine_path.write_text(spine.model_dump_json())
    gap_path = tmp_path / "tx_elements_gap.json"
    _write_gap_artifact(gap_path, [
        {
            "state": "TX", "entity": "Calendar", "element_name": "tx_localCode",
            "discovery": "spine_within_documented_entity",
            "documented_in_source": False,
            "spine_data_type": "String",
            "spine_extension_name": "tx_calendarExtension",
            "rationale": "x",
        },
    ])
    out_path = tmp_path / "tx_scores_gap.json"

    run_gap(state="TX", gap_path=gap_path, spine_path=spine_path, out_path=out_path)
    record = _read(out_path)["scores"][0]
    # Under v18 (issue #84) the +1/+0.5 fork keys on the
    # ``extension_is_necessary`` LLM judgment. On Step 1 / deterministic
    # baseline (this test runs ``run_gap`` without an LLM extract), the
    # judgment is missing → conservative fallback +0.5 necessary_ext +
    # ``extension_necessity_unresolved`` review flag. The Step 3
    # gap-extract refresh re-folds the extracted judgment, so any
    # downstream extension row carrying False then fires +1.
    assert record["adjusted_nachos_score"] is not None
    assert record["adjusted_nachos_score"] >= 0.5
    assert "+0.5 necessary_ext" in (record["nachos_justification"] or "")
    assert "extension_necessity_unresolved" in record["review"]["reasons"]


def test_sidecar_echoes_surfacer_envelope(tmp_path: Path) -> None:
    """Header carries ``gap_source_generated_at`` + ``gap_discovery_counts``.

    Lets consumers (workbook / reviewer comparison) see the source
    artifact's provenance without reopening it.
    """
    spine = _spine(
        entities={"Calendar": EntityEntry(properties={"calendarCode": PropertyInfo()})}
    )
    spine_path = tmp_path / "tx_spine.json"
    spine_path.write_text(spine.model_dump_json())
    gap_path = tmp_path / "tx_elements_gap.json"
    _write_gap_artifact(gap_path, [
        {
            "state": "TX", "entity": "Calendar", "element_name": "calendarCode",
            "discovery": "spine_within_documented_entity",
            "documented_in_source": False,
            "spine_data_type": "String",
            "spine_extension_name": None,
            "rationale": "x",
        },
        {
            "state": "TX", "entity": "Calendar", "element_name": "schoolReference",
            "discovery": "spine_only_full_entity",
            "documented_in_source": False,
            "spine_data_type": "Reference",
            "spine_extension_name": None,
            "rationale": "x",
        },
    ])
    out_path = tmp_path / "tx_scores_gap.json"

    header = run_gap(
        state="TX", gap_path=gap_path, spine_path=spine_path, out_path=out_path
    )

    assert header["gap_source_generated_at"] == "2026-04-29T00:00:00Z"
    assert header["gap_discovery_counts"]["spine_within_documented_entity"] == 1
    assert header["gap_discovery_counts"]["spine_only_full_entity"] == 1
    assert header["step"] == "1"


# ---------------------------------------------------------------------------
# CLI wiring contract — module-level run() must be a plain function so cli.py
# can import + call it directly. Click decoration breaks that contract; the
# regression test mirrors `tests/test_ingest_az.py::TestCliWiring`.
# ---------------------------------------------------------------------------


class TestCliWiring:
    def test_run_is_plain_function(self) -> None:
        from src.score import aggregate_gap

        # A click-wrapped function would carry a `.callback` attribute and
        # be a `click.Command` instance. Plain functions do not.
        assert not hasattr(aggregate_gap.run, "callback"), (
            "aggregate_gap.run must be a plain function — Click wrapper "
            "lives in cli.py per CLAUDE.md."
        )


# ---------------------------------------------------------------------------
# Workbook integration — Spine Gap sheet renders when both gap artifacts
# are present alongside the spine elements artifact.
# ---------------------------------------------------------------------------


def _seed_spine_workbook_fixture(tmp: Path) -> None:
    """Lay down minimum spine-lens elements + gap artifacts under tmp/data/out."""
    from src.models.element import ElementRecord, StateElements

    out = tmp / "data" / "out"
    spine_dir = tmp / "data" / "spine"
    out.mkdir(parents=True, exist_ok=True)
    spine_dir.mkdir(parents=True, exist_ok=True)

    # Spine catalog: Calendar entity (one property + one reference into School).
    spine = _spine(
        entities={
            "School": EntityEntry(
                properties={"schoolId": PropertyInfo(is_identity=True)},
            ),
            "Calendar": EntityEntry(
                properties={"calendarCode": PropertyInfo()},
                references={
                    "schoolReference": ReferenceInfo(
                        entity="School",
                        key_properties={"schoolId": PropertyInfo(is_identity=True)},
                    )
                },
            ),
        }
    )
    (spine_dir / "tx_spine.json").write_text(spine.model_dump_json())

    # One spine-lens elements row so analyst._load_inputs() succeeds.
    elements = StateElements(
        state="TX",
        edfi_version="4.0",
        extracted_at=datetime.now(timezone.utc),
        element_count=1,
        elements=[
            ElementRecord(
                state="TX", edfi_version="4.0", domain="SchoolCalendar",
                entity="Calendar", element_name="calendarCode",
                definition_text="x", source="core", documented=True,
            ),
        ],
    )
    (out / "tx_elements_spine.json").write_text(
        elements.model_dump_json(indent=2), encoding="utf-8"
    )
    # Source-lens elements artifact also expected by _load_inputs in some
    # branches — write a permissive empty one if the spine-lens path needs
    # it. analyst._load_inputs(state, lens="spine") only reads the spine
    # path, so this is informational only but cheap to provide.
    (out / "tx_elements_source.json").write_text(
        elements.model_dump_json(indent=2), encoding="utf-8"
    )

    # Gap artifact + sidecar.
    gap_path = out / "tx_elements_gap.json"
    _write_gap_artifact(gap_path, [
        {
            "state": "TX", "entity": "Calendar", "element_name": "schoolReference",
            "discovery": "spine_within_documented_entity",
            "documented_in_source": False,
            "spine_data_type": "Reference",
            "spine_extension_name": None,
            "rationale": "x",
        },
    ])
    run_gap(
        state="TX",
        gap_path=gap_path,
        spine_path=spine_dir / "tx_spine.json",
        out_path=out / "tx_scores_gap.json",
    )


def test_workbook_renders_spine_gap_sheet(tmp_path: Path, monkeypatch) -> None:
    """End-to-end: a spine-lens workbook gains the "API Model Gaps" sheet
    (formerly "Spine Gap", #174) when both `{state}_elements_gap.json` and
    `{state}_scores_gap.json` are present in the out-dir."""
    import openpyxl

    from src.report import analyst

    _seed_spine_workbook_fixture(tmp_path)

    out = tmp_path / "data" / "out"
    monkeypatch.setattr(analyst, "_OUT_DIR", out)
    monkeypatch.setattr(analyst, "_SPINE_DIR", tmp_path / "data" / "spine")

    [path] = analyst.run(state="TX", out_dir=out, lens="spine")
    wb = openpyxl.load_workbook(path)
    assert "API Model Gaps" in wb.sheetnames
    sheet = wb["API Model Gaps"]
    headers = [c.value for c in sheet[1]]
    assert headers[:7] == [
        "State", "Entity", "Element", "Discovery", "API Model Data Type",
        "Extension Schema", "Implementation Shape",
    ]
    # One data row matching the synthesized gap entry.
    assert sheet.max_row == 2
    assert sheet.cell(row=2, column=2).value == "Calendar"
    assert sheet.cell(row=2, column=3).value == "schoolReference"
    assert sheet.cell(row=2, column=4).value == "spine_within_documented_entity"


def test_workbook_omits_spine_gap_sheet_when_sidecar_missing(
    tmp_path: Path, monkeypatch
) -> None:
    """Without `{state}_scores_gap.json`, the API Model Gaps sheet must not
    render — preserves the existing per-state spine workbook shape unchanged."""
    import openpyxl

    from src.models.element import ElementRecord, StateElements
    from src.report import analyst

    out = tmp_path / "data" / "out"
    spine_dir = tmp_path / "data" / "spine"
    out.mkdir(parents=True, exist_ok=True)
    spine_dir.mkdir(parents=True, exist_ok=True)

    spine = _spine(
        entities={"Calendar": EntityEntry(properties={"calendarCode": PropertyInfo()})}
    )
    (spine_dir / "tx_spine.json").write_text(spine.model_dump_json())
    elements = StateElements(
        state="TX", edfi_version="4.0",
        extracted_at=datetime.now(timezone.utc), element_count=1,
        elements=[
            ElementRecord(
                state="TX", edfi_version="4.0", domain="SchoolCalendar",
                entity="Calendar", element_name="calendarCode",
                definition_text="x", source="core", documented=True,
            ),
        ],
    )
    (out / "tx_elements_spine.json").write_text(elements.model_dump_json())
    (out / "tx_elements_source.json").write_text(elements.model_dump_json())

    monkeypatch.setattr(analyst, "_OUT_DIR", out)
    monkeypatch.setattr(analyst, "_SPINE_DIR", spine_dir)
    [path] = analyst.run(state="TX", out_dir=out, lens="spine")
    wb = openpyxl.load_workbook(path)
    assert "API Model Gaps" not in wb.sheetnames

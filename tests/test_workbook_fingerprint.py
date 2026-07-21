"""R2 workbook-shape fingerprint golden.

Builds a fully deterministic synthetic AZ workbook (fixed timestamps)
for each lens and compares its structural fingerprint — sheet order,
headers, dimensions, widths, freeze/autofilter, tab colors, number
formats, conditional-formatting rules, and a hash over every cell value
— against a committed golden. This is the durable CI form of the
`scripts/diff_workbooks.py` shape-preservation gate: any unintended
presentation change fails here with a reviewable JSON diff.

Regenerate deliberately (e.g. the Option B reshape) with:
    PYTHONPATH=src .venv/bin/python scripts/refresh_workbook_fingerprints.py

The only excluded cell is the value beside the "Workbook generated"
label (stamped at build time).
"""

from __future__ import annotations

import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook

from src.models.edfi_catalog import (
    EdFiCatalog,
    EntityEntry,
    ExtensionEntry,
    PropertyInfo,
    ReferenceInfo,
)
from src.models.element import ElementRecord, StateElements
from src.models.spine import SpineSourceURLs, StateSpine
from src.report import analyst

_GOLDEN_DIR = Path(__file__).resolve().parent / "golden"

_FIXED_DT = datetime(2026, 1, 1, 0, 0, 0, tzinfo=timezone.utc)


def _fixture_spine() -> StateSpine:
    ref = ReferenceInfo(
        entity="School",
        key_properties={"schoolId": PropertyInfo(type="integer", is_identity=True)},
    )
    calendar = EntityEntry(
        description="Calendar",
        domains=["SchoolCalendar"],
        properties={"calendarCode": PropertyInfo(max_length=60)},
        references={"schoolReference": ref},
    )
    school = EntityEntry(description="School", domains=["EducationOrganization"])
    ext = ExtensionEntry(
        extends_entity="Calendar",
        source_prefix="az",
        properties={"az_extra": PropertyInfo()},
    )
    catalog = EdFiCatalog(
        version="4.0.0",
        entity_count=2,
        extension_count=1,
        entities={"Calendar": calendar, "School": school},
        extensions={"CalendarExtension": ext},
    )
    return StateSpine(
        state="AZ",
        edfi_version="4.0.0",
        fetched_at=_FIXED_DT,
        source_urls=SpineSourceURLs(
            resources="https://example/AZ/resources/swagger.json"
        ),
        catalog=catalog,
    )


def _fixture_elements() -> StateElements:
    records = [
        ElementRecord(
            state="AZ",
            edfi_version="4.0",
            domain="SchoolCalendar",
            edfi_domain="SchoolCalendar",
            entity="Calendar",
            element_name="calendarCode",
            data_type="string",
            definition_text="A unique code for the calendar.",
            business_rules_text="Must be unique within the school.",
            source="core",
            source_document="calendar.pdf",
            source_page_or_section="p.3",
            documented=True,
        ),
        ElementRecord(
            state="AZ",
            edfi_version="4.0",
            domain="SchoolCalendar",
            edfi_domain="SchoolCalendar",
            entity="Calendar",
            element_name="az_extra",
            data_type="string",
            definition_text="AZ extension field.",
            source="extension",
            extension_name="az",
            documented=True,
        ),
        ElementRecord(
            state="AZ",
            edfi_version="4.0",
            domain="EducationOrganization",
            edfi_domain="EducationOrganization",
            entity="School",
            element_name="schoolId",
            data_type="integer",
            definition_text="Local school identifier.",
            source="core",
            documented=True,
        ),
        # Swagger-backfill row — blank Source Area, sorts last.
        ElementRecord(
            state="AZ",
            edfi_version="4.0",
            domain="",
            edfi_domain="EducationOrganization",
            entity="School",
            element_name="nameOfInstitution",
            data_type="string",
            definition_text="",
            source="core",
            documented=False,
            documentation_source="swagger",
        ),
    ]
    return StateElements(
        state="AZ",
        edfi_version="4.0",
        extracted_at=_FIXED_DT,
        element_count=len(records),
        elements=records,
    )


# Deliberately divergent from tests/factories.make_score: the fingerprint
# fixture carries extra dimensions + fact_provenance and is parameter-driven;
# its bytes are pinned by the committed goldens, so it must not drift with
# the shared factory.
def _fixture_score(record: ElementRecord, *, tier: int, rule: str,
                   adjusted: float, justification: str) -> dict:
    return {
        "record_key": f"AZ|{record.entity}|{record.element_name}",
        "entity": record.entity,
        "element_name": record.element_name,
        "dimensions": {
            "canonical_name_alignment": {"value": 2, "confidence": "high"},
            "definition_quality": {"value": 2, "confidence": "medium"},
            "semantic_fidelity": {"value": 3, "confidence": "high"},
            "extension_justification": {"value": None, "confidence": "high"},
            "documentation_completeness": {"value": 3, "confidence": "high"},
            "obligation_clarity": {"value": 2, "confidence": "high"},
            "business_logic_complexity": {"value": 1, "confidence": "high"},
            "nachos_score": {
                "value": tier,
                "rule_matched": rule,
                "confidence": "high",
                "inputs_used": {
                    "has_conditional_logic__reconciled": tier >= 1,
                    "has_cross_entity_logic__reconciled": False,
                    "cross_entity_targets": 0,
                },
            },
            "structural_depth": {"value": 2, "confidence": "high"},
            "documentation_style_tier": {
                "value": 1,
                "confidence": "high",
                "inputs_used": {"documentation_style": "conceptual"},
            },
            "documentation_gap": {"value": 0, "confidence": "high"},
        },
        "complexity_score": 1,
        "confidence_composite": "high",
        "fact_provenance": {
            "has_conditional_logic": {
                "value": tier >= 1,
                "spans": ["if enrolled after October 1"] if tier >= 1 else [],
            },
            "semantic_class": {"value": "aligned", "spans": ["span a"]},
        },
        "review": {
            "needs_review": record.source == "extension",
            "reasons": (
                ["low_confidence_dimension:definition_quality"]
                if record.source == "extension"
                else []
            ),
            "route": "SCORING" if record.source == "extension" else None,
        },
        "adjusted_nachos_score": adjusted,
        "in_scope": True,
        "nachos_justification": justification,
    }


def _write_fixture(tmp: Path) -> None:
    out = tmp / "data" / "out"
    spine_dir = tmp / "data" / "spine"
    out.mkdir(parents=True, exist_ok=True)
    spine_dir.mkdir(parents=True, exist_ok=True)
    elements = _fixture_elements()
    for lens in ("source", "spine"):
        (out / f"az_elements_{lens}.json").write_text(
            elements.model_dump_json(indent=2), encoding="utf-8"
        )
    (spine_dir / "az_spine.json").write_text(
        _fixture_spine().model_dump_json(indent=2), encoding="utf-8"
    )
    (out / "az_gap_log.json").write_text(
        json.dumps({
            "state": "AZ",
            "source_coverage": {"matched": 3, "total": 4, "pct": 75.0},
            "spine_coverage": {
                "matched_unique_keys": 3,
                "total_spine_keys": 100,
                "pct": 3.0,
            },
        }),
        encoding="utf-8",
    )
    records = elements.elements
    scores = [
        _fixture_score(records[0], tier=1, rule="tier_1_conditional",
                       adjusted=1.0, justification="tier_1_conditional"),
        _fixture_score(records[1], tier=0, rule="tier_0_none", adjusted=0.5,
                       justification="tier_0_none; +0.5 necessary_ext"),
        _fixture_score(records[2], tier=0, rule="tier_0_none", adjusted=0.0,
                       justification="tier_0_none"),
    ]
    for lens in ("source", "spine"):
        (out / f"az_scores_{lens}.json").write_text(
            json.dumps({
                "state": "AZ",
                "lens": lens,
                "record_count": len(records),
                "scored_count": len(scores),
                "needs_review_count": 1,
                "scores": scores,
            }),
            encoding="utf-8",
        )
        # Recommendations existence gate — analyst.run() only checks the
        # file EXISTS, then regenerates rows from the scores sidecar; the
        # gate brings the Recommendations sheet + Commitment Tracker
        # (Option C) into the fingerprint surface.
        (out / f"az_recommendations_{lens}.json").write_text(
            "{}", encoding="utf-8"
        )


def build_fixture_workbook(tmp: Path, lens: str) -> Path:
    """Build the deterministic AZ workbook for one lens; return its path."""
    _write_fixture(tmp)
    out = tmp / "data" / "out"
    with mock.patch.object(analyst, "_OUT_DIR", out), mock.patch.object(
        analyst, "_SPINE_DIR", tmp / "data" / "spine"
    ):
        paths = analyst.run(state="AZ", out_dir=out, lens=lens)
    return paths[0]


def _timestamp_cells(ws) -> set[tuple[int, int]]:
    cells = set()
    for row in ws.iter_rows():
        for cell in row:
            if cell.value == "Workbook generated":
                cells.add((cell.row, cell.column + 1))
    return cells


def fingerprint_workbook(path: Path) -> dict:
    wb = load_workbook(path)
    fp: dict = {"sheets": {}, "sheet_order": wb.sheetnames}
    for ws in wb.worksheets:
        skip = _timestamp_cells(ws)
        value_hash = hashlib.sha256()
        format_hash = hashlib.sha256()
        for row in ws.iter_rows():
            for cell in row:
                if (cell.row, cell.column) in skip:
                    continue
                if cell.value is not None:
                    value_hash.update(
                        f"{cell.coordinate}={cell.value!r}\n".encode()
                    )
                    format_hash.update(
                        f"{cell.coordinate}|{cell.number_format}|"
                        f"{bool(cell.alignment.wrap_text)}|{cell.font.bold}\n".encode()
                    )
        cf: list = []
        for fmt in ws.conditional_formatting:
            for rule in fmt.rules:
                cf.append(f"{fmt.sqref}|{rule.type}|{rule.operator}")
        fp["sheets"][ws.title] = {
            "headers": [c.value for c in ws[1]],
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "freeze": ws.freeze_panes,
            "autofilter": ws.auto_filter.ref,
            "tab_color": getattr(ws.sheet_properties.tabColor, "rgb", None),
            "widths": {
                k: v.width
                for k, v in sorted(ws.column_dimensions.items())
                if v.width
            },
            "conditional_formatting": sorted(cf),
            "value_sha256": value_hash.hexdigest(),
            "style_sha256": format_hash.hexdigest(),
        }
    return fp


def _dump_drift_details(
    tmp_path: Path, wb_path: Path, golden: dict, actual: dict, lens: str
) -> tuple[Path, list[str]]:
    """Write per-cell values for every drifted sheet to a tmp file.

    The golden stores only ``value_sha256`` per sheet, so a hash mismatch is
    opaque on its own (issue #213 item 4d). This dumps exactly the
    ``coordinate=value`` lines that feed the hash for each sheet whose
    value/style hash (or structural entry) drifted; diff the dump against
    one produced from a clean checkout to see the per-cell change.
    Returns the dump path + the drifted sheet titles.
    """
    g_sheets = golden.get("sheets", {})
    a_sheets = actual.get("sheets", {})
    drifted = sorted(
        title
        for title in a_sheets
        if title not in g_sheets or a_sheets[title] != g_sheets[title]
    )
    removed = sorted(set(g_sheets) - set(a_sheets))
    dump = tmp_path / f"fingerprint_drift_{lens}.txt"
    wb = load_workbook(wb_path)
    lines = [
        f"# per-cell dump of drifted sheets ({lens} lens) — built workbook side",
        f"# golden: {_GOLDEN_DIR / f'workbook_fingerprint_{lens}.json'}",
        f"# sheets missing from built workbook: {removed or 'none'}",
    ]
    for title in drifted:
        lines.append(f"=== {title} ===")
        if title not in g_sheets:
            lines.append("# (sheet absent from golden)")
        else:
            changed = [
                key
                for key in sorted(set(a_sheets[title]) | set(g_sheets[title]))
                if a_sheets[title].get(key) != g_sheets[title].get(key)
            ]
            lines.append(f"# drifted keys: {changed}")
        ws = wb[title]
        skip = _timestamp_cells(ws)
        for row in ws.iter_rows():
            for cell in row:
                if (cell.row, cell.column) in skip or cell.value is None:
                    continue
                lines.append(f"{cell.coordinate}={cell.value!r}")
    dump.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return dump, drifted


class TestWorkbookFingerprint:
    def _assert_matches_golden(self, tmp_path, lens: str):
        golden_path = _GOLDEN_DIR / f"workbook_fingerprint_{lens}.json"
        assert golden_path.exists(), (
            f"missing golden {golden_path} — run "
            "scripts/refresh_workbook_fingerprints.py"
        )
        golden = json.loads(golden_path.read_text(encoding="utf-8"))
        wb_path = build_fixture_workbook(tmp_path, lens)
        actual = fingerprint_workbook(wb_path)
        if actual != golden:
            dump, drifted = _dump_drift_details(
                tmp_path, wb_path, golden, actual, lens
            )
            raise AssertionError(
                "workbook shape drifted from the committed fingerprint — "
                f"drifted sheets: {drifted}. Per-cell values for those "
                f"sheets are dumped at {dump} (diff against a dump from a "
                "clean checkout to see the exact cells). If the change is "
                "intentional, regenerate via "
                "scripts/refresh_workbook_fingerprints.py and review the "
                "JSON diff in the same commit"
            )

    def test_source_lens_fingerprint(self, tmp_path):
        self._assert_matches_golden(tmp_path, "source")

    def test_spine_lens_fingerprint(self, tmp_path):
        self._assert_matches_golden(tmp_path, "spine")

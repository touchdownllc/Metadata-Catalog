"""Focused unit tests for the generic table-sheet renderer
(`src.report.workbook_render.render_table_sheet`) plus the sheet-finish
resolution (`_finish_for` / `_finalize_workbook`).

Everything here runs against tiny in-memory SheetSpec/ColumnSpec fixtures
and a bare openpyxl Workbook() — no analyst fixtures, no disk I/O. The
renderer is row-type agnostic (extractors receive whatever the caller
passes), so most tests use plain dict rows; the Row # injection tests use
real RowContext rows because the renderer reads ``ctx.row_number``.
"""

from __future__ import annotations

import io

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from src.models.element import ElementRecord
from src.report import workbook_render as wr
from src.report.workbook_spec import (
    _DETAILS_FINISH,
    SHEET_SPECS,
    SPINE,
    ColumnSpec,
    RowContext,
    SheetSpec,
)


# --- Tiny in-memory fixtures -------------------------------------------------


def _tiny_columns() -> tuple[ColumnSpec, ...]:
    """3 columns with distinguishable extractors over dict rows; Gamma is
    spine-only so lens filtering is observable."""
    return (
        ColumnSpec("alpha", "Alpha", lambda r: r["alpha"]),
        ColumnSpec("beta", "Beta", lambda r: r["beta"], width=40),
        ColumnSpec("gamma", "Gamma", lambda r: r["gamma"], lens=SPINE),
    )


def _tiny_spec(**overrides) -> SheetSpec:
    base = dict(title="Tiny", columns=_tiny_columns())
    base.update(overrides)
    return SheetSpec(**base)


def _dict_rows(n: int = 2) -> list[dict]:
    return [
        {"alpha": f"a{i}", "beta": f"b{i}", "gamma": f"g{i}"}
        for i in range(1, n + 1)
    ]


def _render(spec: SheetSpec, rows, lens: str = "source", **kw):
    wb = Workbook()
    ws = wb.active
    wr.render_table_sheet(ws, spec, rows, lens, **kw)
    return ws


def _header_col(ws) -> dict[str, int]:
    return {
        ws.cell(row=1, column=c).value: c
        for c in range(1, ws.max_column + 1)
    }


def _ctx(element: str, row_number: int | None = None) -> RowContext:
    record = ElementRecord(
        state="AZ",
        edfi_version="4.0",
        domain="SchoolCalendar",
        entity="Calendar",
        element_name=element,
        definition_text="doc",
        source="core",
        documented=True,
    )
    return RowContext(
        state="AZ",
        record=record,
        is_extension=False,
        edfi_domain="SchoolCalendar",
        row_number=row_number,
    )


_CTX_COLUMNS: tuple[ColumnSpec, ...] = (
    ColumnSpec("state", "State", lambda c: c.state),
    ColumnSpec("element", "Data Element", lambda c: c.record.element_name),
)


class TestRenderTableSheet:
    def test_header_row_equals_spec_labels(self):
        ws = _render(_tiny_spec(), _dict_rows(), lens="spine")
        assert [c.value for c in ws[1]] == ["Alpha", "Beta", "Gamma"]

    def test_header_font_bold_and_grey_fill(self):
        ws = _render(_tiny_spec(), _dict_rows(), lens="spine")
        for cell in ws[1]:
            assert cell.font.bold is True
            assert cell.fill.fill_type == "solid"
            assert cell.fill.start_color.rgb == "FFDDDDDD"

    def test_values_land_under_their_header_by_name(self):
        ws = _render(_tiny_spec(), _dict_rows(2), lens="spine")
        col = _header_col(ws)
        assert ws.cell(row=2, column=col["Alpha"]).value == "a1"
        assert ws.cell(row=2, column=col["Beta"]).value == "b1"
        assert ws.cell(row=2, column=col["Gamma"]).value == "g1"
        assert ws.cell(row=3, column=col["Alpha"]).value == "a2"
        assert ws.cell(row=3, column=col["Gamma"]).value == "g2"
        # Header + 2 data rows, nothing extra.
        assert ws.max_row == 3

    def test_lens_filtering_drops_spine_only_column(self):
        ws = _render(_tiny_spec(), _dict_rows(1), lens="source")
        headers = [c.value for c in ws[1]]
        assert headers == ["Alpha", "Beta"]
        assert "Gamma" not in headers
        # Values pack into the filtered column list — no gap where the
        # spine-only column would have been.
        assert ws.max_column == 2
        col = _header_col(ws)
        assert ws.cell(row=2, column=col["Alpha"]).value == "a1"
        assert ws.cell(row=2, column=col["Beta"]).value == "b1"

    def test_title_defaults_to_spec_title(self):
        ws = _render(_tiny_spec(), [], lens="source")
        assert ws.title == "Tiny"

    def test_title_override_parameter(self):
        # The spine-lens "{ST} — Documented only" family reuses the
        # Reviewer View spec under per-state titles.
        ws = _render(
            _tiny_spec(), [], lens="source", title="AZ — Documented only"
        )
        assert ws.title == "AZ — Documented only"

    def test_per_column_width_and_sheet_default_width(self):
        ws = _render(_tiny_spec(default_width=15), _dict_rows(1), lens="spine")
        col = _header_col(ws)
        # Beta carries an explicit width; the others fall back to the
        # sheet default.
        assert ws.column_dimensions[get_column_letter(col["Beta"])].width == 40
        assert ws.column_dimensions[get_column_letter(col["Alpha"])].width == 15
        assert ws.column_dimensions[get_column_letter(col["Gamma"])].width == 15

    def test_wrap_column_cells_wrap_and_top_align(self):
        columns = (
            ColumnSpec("plain", "Plain", lambda r: r["alpha"]),
            ColumnSpec("prose", "Prose", lambda r: r["beta"], wrap=True),
        )
        ws = _render(
            SheetSpec(title="Wrap", columns=columns), _dict_rows(2), lens="source"
        )
        col = _header_col(ws)
        for row in (2, 3):
            wrapped = ws.cell(row=row, column=col["Prose"]).alignment
            assert wrapped.wrap_text is True
            assert wrapped.vertical == "top"
            plain = ws.cell(row=row, column=col["Plain"]).alignment
            assert not plain.wrap_text

    def test_formula_guard_leading_equals_round_trips_as_text(self):
        # WI-style business-rules prose starts with "=== …" — without the
        # guard openpyxl writes it as a formula and Excel flags the file
        # as corrupt.
        columns = (ColumnSpec("t", "Text", lambda r: r["alpha"]),)
        ws = _render(
            SheetSpec(title="Guard", columns=columns),
            [{"alpha": "=== Element-specific rules ==="}],
            lens="source",
        )
        cell = ws.cell(row=2, column=1)
        assert cell.data_type == "s"
        assert cell.value == "=== Element-specific rules ==="
        # Save/reload round-trip stays text.
        buf = io.BytesIO()
        ws.parent.save(buf)
        reloaded = load_workbook(io.BytesIO(buf.getvalue()))["Guard"]
        again = reloaded.cell(row=2, column=1)
        assert again.data_type == "s"
        assert again.value == "=== Element-specific rules ==="

    def test_columns_factory_spec_renders_via_factory(self):
        lenses_seen: list[str] = []

        def factory(lens: str) -> tuple[ColumnSpec, ...]:
            lenses_seen.append(lens)
            cols = [ColumnSpec("a", f"Alpha ({lens})", lambda r: r["alpha"])]
            if lens == "spine":
                cols.append(ColumnSpec("g", "Gamma", lambda r: r["gamma"]))
            return tuple(cols)

        spec = SheetSpec(title="Factory", columns_factory=factory)
        ws_spine = _render(spec, _dict_rows(1), lens="spine")
        assert [c.value for c in ws_spine[1]] == ["Alpha (spine)", "Gamma"]
        assert ws_spine.cell(row=2, column=2).value == "g1"
        ws_source = _render(spec, _dict_rows(1), lens="source")
        assert [c.value for c in ws_source[1]] == ["Alpha (source)"]
        assert set(lenses_seen) == {"spine", "source"}


class TestRowNumberInjection:
    def test_row_number_header_first_and_values_from_context(self):
        spec = SheetSpec(
            title="Ctx", columns=_CTX_COLUMNS, row_number_column=True
        )
        rows = [_ctx("calendarCode", row_number=5), _ctx("schoolId", row_number=9)]
        ws = _render(spec, rows, lens="source")
        headers = [c.value for c in ws[1]]
        assert headers[0] == "Row #"
        assert headers[1:] == ["State", "Data Element"]
        # The injected leading cell carries ctx.row_number verbatim; the
        # spec columns shift right by one.
        assert ws.cell(row=2, column=1).value == 5
        assert ws.cell(row=3, column=1).value == 9
        col = _header_col(ws)
        assert ws.cell(row=2, column=col["Data Element"]).value == "calendarCode"
        assert ws.cell(row=3, column=col["Data Element"]).value == "schoolId"

    def test_row_number_column_width_is_8(self):
        spec = SheetSpec(
            title="Ctx", columns=_CTX_COLUMNS, row_number_column=True,
            default_width=30,
        )
        ws = _render(spec, [_ctx("calendarCode", row_number=1)], lens="source")
        # Compact Row # address column; spec columns get the default.
        assert ws.column_dimensions["A"].width == 8
        assert ws.column_dimensions["B"].width == 30


class TestFinishResolution:
    def test_unknown_title_raises_keyerror(self):
        # The build-time non-drift guard: an unregistered sheet cannot
        # ship without a guide + finish.
        with pytest.raises(KeyError):
            wr._finish_for("Some Future Sheet")

    def test_documented_only_suffix_resolves_to_shared_details_finish(self):
        # Identity, not just equality — the per-state family shares ONE
        # finish object with the Details spec.
        assert wr._finish_for("AZ — Documented only") is _DETAILS_FINISH
        assert wr._finish_for("WI — Documented only") is _DETAILS_FINISH
        assert SHEET_SPECS["Details"].finish is _DETAILS_FINISH

    def test_finalize_applies_tab_color_from_group(self):
        wb = Workbook()
        wb.active.title = "Readme"          # orient
        wb.create_sheet("Details")          # work
        wb.create_sheet("Audit Trail")      # audit
        wr._finalize_workbook(wb, "source")
        assert wb["Readme"].sheet_properties.tabColor.rgb == "FFBDD7EE"
        assert wb["Details"].sheet_properties.tabColor.rgb == "FFC6E0B4"
        assert wb["Audit Trail"].sheet_properties.tabColor.rgb == "FFD9D9D9"


class TestAnalystInputBandStyling:
    def test_analyst_input_headers_green_with_comment_on_first(self):
        """Option B band styling: analyst-input headers get the light-green
        fill (FFC6E0B4) and the FIRST band header carries the openpyxl
        comment explaining the band; machine columns keep the grey fill."""
        columns = (
            ColumnSpec("machine", "Machine", lambda r: r["alpha"]),
            ColumnSpec(
                "note_a", "Analyst A", lambda r: None, role="analyst_input"
            ),
            ColumnSpec(
                "note_b", "Analyst B", lambda r: None, role="analyst_input"
            ),
        )
        ws = _render(
            SheetSpec(title="Band", columns=columns), _dict_rows(1),
            lens="source",
        )
        col = _header_col(ws)
        machine = ws.cell(row=1, column=col["Machine"])
        assert machine.fill.start_color.rgb == "FFDDDDDD"
        assert machine.comment is None
        first = ws.cell(row=1, column=col["Analyst A"])
        second = ws.cell(row=1, column=col["Analyst B"])
        for cell in (first, second):
            assert cell.fill.fill_type == "solid"
            assert cell.fill.start_color.rgb == "FFC6E0B4"
        # Comment on the FIRST band header only.
        assert first.comment is not None
        assert "pipeline never writes" in first.comment.text
        assert second.comment is None
        # Data cells in the band render blank (extractors return None).
        assert ws.cell(row=2, column=col["Analyst A"]).value is None

    def test_row_number_offset_does_not_shift_band_styling(self):
        """With the injected Row # column the role list shifts by one —
        the green fill must land on the analyst_input headers, not their
        neighbors."""
        columns = (
            ColumnSpec("state", "State", lambda c: c.state),
            ColumnSpec(
                "note", "Analyst Note", lambda c: None, role="analyst_input"
            ),
        )
        spec = SheetSpec(title="Band2", columns=columns, row_number_column=True)
        ws = _render(spec, [_ctx("calendarCode", row_number=1)], lens="source")
        col = _header_col(ws)
        assert ws.cell(row=1, column=col["Row #"]).fill.start_color.rgb == "FFDDDDDD"
        assert ws.cell(row=1, column=col["State"]).fill.start_color.rgb == "FFDDDDDD"
        note = ws.cell(row=1, column=col["Analyst Note"])
        assert note.fill.start_color.rgb == "FFC6E0B4"
        assert note.comment is not None
